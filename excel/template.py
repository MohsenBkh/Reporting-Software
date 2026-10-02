# -*- coding: utf-8 -*-
"""قالب Excel استاندارد نرم‌افزار — تولید و خواندن (بخش ۱۳ و ۲۸ سند)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="B Nazanin", size=11)
THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RTL_SHEET_VIEW = {"rightToLeft": True}

PROJECT_KEYS = [
    ("عنوان پروژه", "name", ""),
    ("شماره گزارش", "report_number", ""),
    ("تاریخ گزارش", "date_jalali", ""),
    ("نام متقاضی", "applicant_name", ""),
    ("نوع درخواست", "request_type", "افزایش قدرت / تأمین برق جدید"),
    ("توان فعلی (kW)", "existing_power_kw", "فقط برای افزایش قدرت"),
    ("توان جدید/درخواستی (kW)", "requested_power_kw", ""),
    ("نام پست فوق توزیع", "substation", ""),
    ("نام امور", "office", ""),
    ("نام کارشناس", "expert_name", ""),
    ("مانور/بازآرایی (بله/خیر)", "maneuver_enabled", ""),
    ("فیدر مبدأ مانور", "maneuver_source", ""),
    ("فیدر مقصد مانور", "maneuver_target", ""),
    ("بار منتقل‌شده (MW)", "maneuver_mw", "تخمینی"),
    ("تاریخ مانور", "maneuver_date", "در صورت مشخص بودن"),
    ("توضیحات مانور", "maneuver_note", "عدم قطعیت‌ها و توضیح کارشناس"),
    ("توضیح موقعیت محل", "location_note", "اختیاری"),
    ("فاصله تا نزدیک‌ترین تیر (متر)", "location_distance_m", "اختیاری"),
    ("پیشنهاد نوع کابل", "cable_suggestion", "اختیاری"),
]

FEEDER_HEADERS = ["نام فیدر", "نام پست", "پیک بار (MW)", "سال پیک", "ضریب توان",
                  "پیک جریان (A)", "حداکثر جریان مجاز (A)", "معیار بارگذاری (MW)", "توضیحات"]

POWERFLOW_HEADERS = ["نام فیدر", "وضعیت", "جریان ابتدای فیدر (A)", "تلفات (kW)",
                     "حداقل ولتاژ فیدر (p.u.)", "ولتاژ باس متقاضی (p.u.)"]

PROFILE_HEADERS = ["Date", "Feeder", "P_MW", "Q_MVAR"]

FORECAST_HEADERS = ["نام فیدر", "سال", "پیک پیش‌بینی‌شده (MW)"]


def _style_header_row(ws, row: int, count: int) -> None:
    for c in range(1, count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def create_template(path: str | Path) -> Path:
    """قالب استاندارد ورودی را می‌سازد."""
    wb = Workbook()

    # --- شیت پروژه ---
    ws = wb.active
    ws.title = "پروژه"
    ws.sheet_view.rightToLeft = True
    ws.append(["فیلد", "مقدار", "توضیح"])
    _style_header_row(ws, 1, 3)
    for label, _key, note in PROJECT_KEYS:
        ws.append([label, "", note])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 36

    # --- شیت فیدرها ---
    ws2 = wb.create_sheet("فیدرها")
    ws2.sheet_view.rightToLeft = True
    ws2.append(FEEDER_HEADERS)
    _style_header_row(ws2, 1, len(FEEDER_HEADERS))
    for i, h in enumerate(FEEDER_HEADERS, 1):
        ws2.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws2.append(["فیدر 1 اردبیل", "کمی‌آباد", 2.75, 1404, 0.93, 85, 400, 7, ""])

    # --- شیت پخش بار ---
    ws3 = wb.create_sheet("پخش بار")
    ws3.sheet_view.rightToLeft = True
    ws3.append(POWERFLOW_HEADERS)
    _style_header_row(ws3, 1, len(POWERFLOW_HEADERS))
    for i, h in enumerate(POWERFLOW_HEADERS, 1):
        ws3.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws3.append(["فیدر 1 اردبیل", "قبل", 201, 452, 0.835, ""])
    ws3.append(["فیدر 1 اردبیل", "بعد", 204, 459, 0.835, 0.954])

    # --- شیت پروفیل بار ---
    ws4 = wb.create_sheet("پروفیل بار")
    ws4.append(PROFILE_HEADERS)
    _style_header_row(ws4, 1, len(PROFILE_HEADERS))
    for i, h in enumerate(PROFILE_HEADERS, 1):
        ws4.column_dimensions[get_column_letter(i)].width = 14
    ws4.append(["1404/05/02", "فیدر 1 اردبیل", 2.5, 0.9])

    # --- شیت پیش‌بینی (اختیاری/دستی) ---
    ws5 = wb.create_sheet("پیش‌بینی")
    ws5.sheet_view.rightToLeft = True
    ws5.append(FORECAST_HEADERS)
    _style_header_row(ws5, 1, len(FORECAST_HEADERS))
    for i, h in enumerate(FORECAST_HEADERS, 1):
        ws5.column_dimensions[get_column_letter(i)].width = 18

    wb.save(path)
    return Path(path)


def _norm_header(value) -> str:
    return (str(value or "").strip().lower()
            .replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
            .replace("_", " ").replace("(", "").replace(")", "")
            .replace(".", "").replace("  ", " ").strip())


class ExcelData:
    """داده خام خوانده‌شده از قالب."""
    def __init__(self) -> None:
        self.project_fields: dict[str, str] = {}
        self.feeders: list[dict] = []
        self.powerflow: list[dict] = []
        self.profile_rows: list[dict] = []
        self.forecast_rows: list[dict] = []


def read_workbook(path: str | Path) -> tuple[ExcelData, list[str]]:
    """خواندن قالب + اعتبارسنجی ساختاری. خروجی: (داده، خطاها)."""
    data = ExcelData()
    errors: list[str] = []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return data, [f"باز کردن فایل Excel ناموفق بود: {exc}"]

    # --- پروژه: key-value ---
    if "پروژه" in wb.sheetnames:
        ws = wb["پروژه"]
        key_map = {_norm_header(label): key for label, key, _ in PROJECT_KEYS}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            key = key_map.get(_norm_header(row[0]))
            if key and len(row) > 1 and row[1] not in (None, ""):
                data.project_fields[key] = str(row[1]).strip()
    else:
        errors.append("شیت «پروژه» پیدا نشد.")

    # --- فیدرها ---
    if "فیدرها" in wb.sheetnames:
        ws = wb["فیدرها"]
        headers = [_norm_header(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = {h: i for i, h in enumerate(headers)}
        if "نام فیدر" not in headers:
            errors.append("در شیت «فیدرها» ستون «نام فیدر» پیدا نشد.")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[col["نام فیدر"]]:
                continue
            rec = {}
            for h in headers:
                if h and col.get(h) is not None and col[h] < len(row):
                    rec[h] = row[col[h]]
            data.feeders.append(rec)
        if not data.feeders and not errors:
            errors.append("در شیت «فیدرها» هیچ فیدری وارد نشده است.")
    else:
        errors.append("شیت «فیدرها» پیدا نشد.")

    # --- پخش بار ---
    if "پخش بار" in wb.sheetnames:
        ws = wb["پخش بار"]
        headers = [_norm_header(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = {h: i for i, h in enumerate(headers)}
        if "نام فیدر" not in headers or "وضعیت" not in headers:
            errors.append("در شیت «پخش بار» ستون «نام فیدر» یا «وضعیت» پیدا نشد.")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            rec = {h: (row[col[h]] if col.get(h) is not None and col[h] < len(row) else None)
                   for h in headers if h}
            name = rec.get("نام فیدر")
            if name:
                data.powerflow.append(rec)
    else:
        errors.append("شیت «پخش بار» پیدا نشد.")

    # --- پروفیل بار (اختیاری) ---
    if "پروفیل بار" in wb.sheetnames:
        ws = wb["پروفیل بار"]
        headers = [_norm_header(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = {h: i for i, h in enumerate(headers)}
        if "date" not in headers or "p mw" not in headers:
            errors.append("در شیت «پروفیل بار» ستون‌های Date یا P_MW پیدا نشد.")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            rec = {}
            for std, aliases in (("date", ["date", "تاریخ"]),
                                 ("feeder", ["feeder", "فیدر"]),
                                 ("p_mw", ["p mw", "p", "توان اکتیو", "mw"]),
                                 ("q_mvar", ["q mvar", "q", "توان راکتیو"])):
                for a in aliases:
                    if a in col and col[a] < len(row):
                        rec[std] = row[col[a]]
                        break
            if rec.get("date") is not None and rec.get("p_mw") is not None:
                data.profile_rows.append(rec)

    # --- پیش‌بینی دستی (اختیاری) ---
    if "پیش‌بینی" in wb.sheetnames:
        ws = wb["پیش‌بینی"]
        headers = [_norm_header(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = {h: i for i, h in enumerate(headers)}
        k_name, k_year = _norm_header("نام فیدر"), _norm_header("سال")
        k_val = _norm_header("پیک پیش‌بینی‌شده (MW)")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            name = row[col[k_name]] if k_name in col and col[k_name] < len(row) else None
            year = row[col[k_year]] if k_year in col and col[k_year] < len(row) else None
            val = row[col[k_val]] if k_val in col and col[k_val] < len(row) else None
            if name and year is not None and val is not None:
                data.forecast_rows.append({"feeder": str(name).strip(),
                                           "year": year, "value": val})
    wb.close()
    return data, errors
