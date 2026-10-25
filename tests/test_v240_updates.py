# -*- coding: utf-8 -*-
"""آزمون‌های v2.4.0 —

۱) رفع اشکال ورود بسته IDMT: فرم‌های باز نباید دادهٔ واردشده را بازنویسی کنند
   (ریکلوزر تنها، ریکلوزر + یک یا چند سکشنالایزر، روی هر صفحه/تب).
۲) قابلیت فیوز سیوینگ (Fuse saving) برای تنظیمات ریکلوزر: مدل، فرم، جدول گزارش
   و انتقال خودکار از بخش «فیوز» بسته IDMT.
۳) مطالعه جدید «توسعه و اصلاح فیدر»: مدل، برنامه گزارش، فصل‌ها، اعتبارسنجی،
   عنوان جلد/سربرگ و خروجی Word.
"""
from __future__ import annotations

import copy
import json
import os
import tempfile
from pathlib import Path

import pytest

os.environ.setdefault("REPORTFORGE_HOME", tempfile.mkdtemp(prefix="rf_home_v240_"))

from app import pyside_compat  # noqa: F401
from PySide6.QtWidgets import QApplication, QTableWidgetItem

QApplication.instance() or QApplication([])

from app.core.idmt_bundle import apply_bundle_to_project, import_idmt_file
from app.core.models import (FeederDevRow, Project, project_from_dict, to_dict)
from app.core.protection_models import ProtectionSettings
from app.core.report_types import (booklet_title, header_title, is_implemented)
from app.core.settings import AppSettings
from app.core.studies import StudyKind, StudySelection, selection_from_legacy
from app.core.validation import ERROR, validate_project
from app.core.workspace import Workspace
from app.report.composer import ReportComposer, plan_reports
from test_v220_idmt_link import make_bundle


# ---------------------------------------------------------------------------
# ابزار مشترک: پنجره مهندسی با پروژه باز
# ---------------------------------------------------------------------------
def _window(tmp_path, *, recloser=True, sectionalizer=False, feeder_dev=False):
    from app.ui.shell.engineering_window import EngineeringWindow
    w = EngineeringWindow()
    w.manager.new_project(tmp_path / "proj")
    w._shell_ready = True
    ws = w._ws()
    ws.selection.heavy = False
    ws.selection.recloser = recloser
    ws.selection.sectionalizer = sectionalizer
    ws.selection.feeder_dev = feeder_dev
    return w


def _bundle_file(tmp_path, bundle) -> Path:
    f = tmp_path / "IDMT08-ReportForge.json"
    f.write_text(json.dumps(bundle, ensure_ascii=False), encoding="utf-8")
    return f


def _recloser_only_bundle():
    """بسته واقعی حالت «ریکلوزر تنها» — بدون هیچ سکشنالایزری."""
    b = copy.deepcopy(make_bundle())
    for mode in ("phase", "earth"):
        b["resolved"][mode]["sections"] = []
    b["resolved"]["cycle"]["sections"] = []
    return b


def _multi_section_bundle(n=3):
    b = copy.deepcopy(make_bundle())
    for mode in ("phase", "earth"):
        base = b["resolved"][mode]["sections"][0]
        b["resolved"][mode]["sections"] = [
            {**copy.deepcopy(base), "id": f"S{i+1}", "name": f"سکشنالایزر {i+1}"}
            for i in range(n)]
    cyc = b["resolved"]["cycle"]["sections"][0]
    b["resolved"]["cycle"]["sections"] = [
        {**copy.deepcopy(cyc), "id": f"S{i+1}"} for i in range(n)]
    return b


# ---------------------------------------------------------------------------
# ۱) رفع اشکال ورود IDMT — فرم باز نباید داده واردشده را پاک کند
# ---------------------------------------------------------------------------
def test_import_recloser_only_while_settings_tab_open(tmp_path):
    """گزارش کاربر: مطالعه تک‌ریکلوزره؛ تنظیمات واردشده پس از ناوبری می‌پرید."""
    w = _window(tmp_path)
    w._open_tab("recloser")
    w.pages["recloser"].refresh()
    w.pages["recloser"].show_tab("settings")        # فرم خالی روی صفحه است
    res = w._apply_idmt_import(_bundle_file(tmp_path, _recloser_only_bundle()))
    w.navigate("recloser")                          # قبل از رفع اشکال، همین‌جا پاک می‌شد
    rec = w.manager.project.recloser
    assert rec.settings.f51.pickup_a == 190.0
    assert rec.settings.shots == 2
    assert rec.settings.reclaim_time_s == 300
    assert rec.upstream_settings.f51.pickup_a == 200.0
    assert rec.feeder_name == "فیدر ۴ خلخال"
    assert res.applied


def test_import_recloser_only_while_device_tab_open(tmp_path):
    w = _window(tmp_path)
    w._open_tab("recloser")
    w.pages["recloser"].refresh()
    w.pages["recloser"].show_tab("device")
    w._apply_idmt_import(_bundle_file(tmp_path, _recloser_only_bundle()))
    w.navigate("recloser")
    rec = w.manager.project.recloser
    assert rec.feeder_name == "فیدر ۴ خلخال"
    assert rec.fault_current_ka == 1.0              # 1000 A → 1 kA
    assert rec.load_current_a == 150
    # بسته بدون سکشنالایزر: مدل سکشنالایزر دست‌نخورده می‌ماند
    assert w.manager.project.sectionalizer.settings.is_empty()


def test_import_with_one_sectionalizer_while_secz_page_open(tmp_path):
    w = _window(tmp_path, sectionalizer=True)
    w._open_tab("sectionalizer")
    w.pages["sectionalizer"].refresh()
    w.pages["sectionalizer"].show_tab("settings")
    w._apply_idmt_import(_bundle_file(tmp_path, make_bundle()))
    w.navigate("sectionalizer")
    secz = w.manager.project.sectionalizer
    assert secz.settings.f51.pickup_a == 180.0
    assert secz.settings.counter == 2
    assert w.manager.project.recloser.settings.f51.pickup_a == 190.0


def test_import_with_three_sectionalizers_keeps_first_and_warns(tmp_path):
    w = _window(tmp_path, sectionalizer=True)
    w._open_tab("sectionalizer")
    w.pages["sectionalizer"].refresh()
    res = w._apply_idmt_import(_bundle_file(tmp_path, _multi_section_bundle(3)))
    w.navigate("sectionalizer")
    assert any("۳" in wn or "3" in wn for wn in res.warnings)
    secz = w.manager.project.sectionalizer
    assert secz.name == "سکشنالایزر 1"
    assert secz.settings.f51.pickup_a == 180.0
    assert w.manager.project.recloser.settings.f51.pickup_a == 190.0


def test_import_while_protnet_page_open_keeps_common_network(tmp_path):
    w = _window(tmp_path)
    w._open_tab("protnet")
    w.prot_network_page.refresh()
    w._apply_idmt_import(_bundle_file(tmp_path, _recloser_only_bundle()))
    w.navigate("protnet")
    ws = w._ws()
    fd = ws.common_network.feeder("فیدر ۴ خلخال")
    assert fd is not None and fd.voltage_kv == 20
    assert "فیدر ۴ خلخال" in ws.scope("recloser")


def test_import_via_import_idmt_file_direct(tmp_path):
    """مسیر بدون UI (idmt_bundle) هم برای ریکلوزر تنها درست کار می‌کند."""
    from app.core.project_manager import ProjectManager
    mgr = ProjectManager(AppSettings())
    mgr.new_project(tmp_path / "p2")
    res = import_idmt_file(mgr, _bundle_file(tmp_path, _recloser_only_bundle()))
    assert mgr.project.recloser.settings.f51.pickup_a == 190.0
    assert len(res.images) == 5


# ---------------------------------------------------------------------------
# ۲) فیوز سیوینگ
# ---------------------------------------------------------------------------
def test_fuse_saving_round_trip_and_is_empty():
    s = ProtectionSettings()
    assert s.is_empty()
    s.fuse_saving = True
    s.fuse_saving_note = "شات‌های ۱ و ۲ روی منحنی سریع"
    assert not s.is_empty()
    s2 = ProtectionSettings.from_dict(s.to_dict())
    assert s2.fuse_saving is True
    assert s2.fuse_saving_note == "شات‌های ۱ و ۲ روی منحنی سریع"


def test_fuse_saving_editor_load_store():
    from app.ui.shell.settings_editor import ProtectionSettingsEditor
    ed = ProtectionSettingsEditor(reclosing=True, counter=False)
    s = ProtectionSettings(fuse_saving=False, fuse_saving_note="غیرفعال در این فیدر")
    ed.load(s)
    out = ProtectionSettings()
    ed.store(out)
    assert out.fuse_saving is False
    assert out.fuse_saving_note == "غیرفعال در این فیدر"
    ed.fuse_saving.setCurrentIndex(0)               # «ثبت نشده»
    ed.store(out)
    assert out.fuse_saving is None
    assert ed.focus_field("fuse_saving")
    # ادیتور سکشنالایزر (بدون reclosing) فیلد فیوز سیوینگ ندارد
    ed2 = ProtectionSettingsEditor(reclosing=False, counter=True)
    assert ed2.fuse_saving is None


def test_fuse_saving_row_in_report_table():
    p = Project()
    p.report_type = "recloser"
    p.recloser.settings.f51.pickup_a = 190
    p.recloser.settings.fuse_saving = True
    p.recloser.settings.fuse_saving_note = "فاز: فیوز K 100A"
    ws = Workspace()
    ws.selection.recloser = True
    rep = ReportComposer().compose(p, ws, [StudyKind.RECLOSER])
    cells = [cell for s in rep.sections for b in s.blocks if hasattr(b, "body_rows")
             for row in b.body_rows for cell in row]
    assert "فیوز سیوینگ (Fuse saving)" in cells
    assert "فعال" in cells
    assert "فاز: فیوز K 100A" in cells
    # بدون ثبت → سطر درج نمی‌شود
    p.recloser.settings.fuse_saving = None
    p.recloser.settings.fuse_saving_note = ""
    rep2 = ReportComposer().compose(p, ws, [StudyKind.RECLOSER])
    cells2 = [cell for s in rep2.sections for b in s.blocks if hasattr(b, "body_rows")
              for row in b.body_rows for cell in row]
    assert "فیوز سیوینگ (Fuse saving)" not in cells2


def test_fuse_saving_from_idmt_bundle():
    b = copy.deepcopy(make_bundle())
    b["project"]["fuse"] = {"phase": {"enabled": True, "type": "K", "rating": 100, "k": 0.75},
                            "earth": {"enabled": False, "type": "K", "rating": 100, "k": 0.75}}
    p = Project()
    res = apply_bundle_to_project(p, b)
    assert p.recloser.settings.fuse_saving is True
    assert "K 100A" in p.recloser.settings.fuse_saving_note
    assert any("فیوز سیوینگ" in a for a in res.applied)
    # بسته بدون بخش فیوز → مقدار دستی قبلی حفظ می‌شود
    p2 = Project()
    p2.recloser.settings.fuse_saving = True
    p2.recloser.settings.fuse_saving_note = "دستی"
    apply_bundle_to_project(p2, make_bundle())
    assert p2.recloser.settings.fuse_saving is True
    assert p2.recloser.settings.fuse_saving_note == "دستی"


# ---------------------------------------------------------------------------
# ۳) مطالعه «توسعه و اصلاح فیدر»
# ---------------------------------------------------------------------------
def _fd_project() -> Project:
    p = Project()
    p.name = "اصلاح خط مسیر اوچ‌بلاغ به سمت کوراماللو"
    p.report_number = "FD-15-001"
    p.date_jalali = "1405/08/03"
    p.report_type = "feeder_dev"
    fd = p.feeder_dev
    fd.plan_kind = "اصلاح خط موجود"
    fd.substation = "پست فوق توزیع خلخال"
    fd.feeders = [FeederDevRow(name="پست خلخال — فیدر ۴", voltage_kv=20, length_km=45.2,
                               peak_load_mw=3.2, min_load_mw=0.9, loading_pct=62,
                               power_factor=0.88)]
    fd.existing_description = "فیدر ۴ خلخال با افت ولتاژ انتهای خط روبه‌رو است."
    fd.proposal_description = "احداث خط جدید به طول ۲٫۵ کیلومتر و برچیدن خط فرسوده."
    fd.new_line_length_km = 2.5
    fd.conductor_type = "هاینا ۱۲۶"
    fd.removed_line_length_km = 4.1
    fd.min_voltage_kv_before = 18.8
    fd.min_voltage_kv_after = 19.6
    fd.losses_kw_before = 210
    fd.losses_kw_after = 150
    return p


def test_feeder_dev_selection_and_plans():
    sel = StudySelection(feeder_dev=True)
    assert sel.kinds() == [StudyKind.FEEDER_DEV]
    assert not sel.protection_kinds()
    assert StudySelection.from_dict(sel.to_dict()).feeder_dev is True
    assert selection_from_legacy("feeder_dev").feeder_dev is True
    plans = plan_reports(sel)
    assert [(p.key, p.filename_stem) for p in plans] == \
        [("feeder_dev", "Feeder Development Report")]
    # ترکیب با مطالعات دیگر: گزارش توسعه فیدر جدا می‌ماند
    plans2 = plan_reports(StudySelection(recloser=True, feeder_dev=True))
    assert [p.key for p in plans2] == ["recloser", "feeder_dev"]


def test_feeder_dev_model_round_trip():
    p = _fd_project()
    p2 = project_from_dict(to_dict(p))
    fd = p2.feeder_dev
    assert fd.plan_kind == "اصلاح خط موجود"
    assert isinstance(fd.feeders[0], FeederDevRow)
    assert fd.feeders[0].length_km == 45.2
    assert fd.min_voltage_kv_after == 19.6
    assert fd.has_results()


def test_feeder_dev_compose_sections_and_tables():
    p = _fd_project()
    ws = Workspace()
    ws.selection.feeder_dev = True
    rep = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    keys = [s.key for s in rep.sections]
    assert keys == ["fd_intro", "fd_network", "fd_before", "fd_plan",
                    "fd_after", "fd_conclusion"]
    # هیچ بخش حفاظتی (prot_*) در گزارش توسعه فیدر نمی‌آید
    assert not [k for k in keys if k.startswith("prot_")]
    captions = [b.caption for s in rep.sections for b in s.blocks
                if hasattr(b, "caption") and hasattr(b, "body_rows")]
    assert any("مشخصات فیدرهای محدوده مطالعه" in c for c in captions)
    assert any("مشخصات طرح پیشنهادی" in c for c in captions)
    assert any("مقایسه نتایج پخش بار قبل و بعد" in c for c in captions)
    # مقدمه پیش‌فرض مرجع‌های نمونه را دارد
    intro_text = " ".join(b.text for b in rep.sections[0].paragraphs())
    assert "WI150900" in intro_text and "TAV111-03/00" in intro_text
    # جمع‌بندی خودکار مقادیر قبل/بعد را بازتاب می‌دهد
    concl = " ".join(b.text for b in rep.sections[-1].paragraphs())
    assert "18.8" in concl and "19.6" in concl


def test_feeder_dev_empty_info_produces_minimal_report():
    p = Project()
    p.report_type = "feeder_dev"
    ws = Workspace()
    ws.selection.feeder_dev = True
    rep = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    keys = [s.key for s in rep.sections]
    assert "fd_intro" in keys                    # مقدمه پیش‌فرض همیشه هست
    assert "fd_network" not in keys              # بخش بدون داده حذف می‌شود


def test_feeder_dev_protection_report_unaffected():
    """گزارش ریکلوزر نباید بخش fd_* بگیرد؛ بخش‌های عمومی حفاظتی سر جایشان‌اند."""
    p = Project()
    p.report_type = "recloser"
    p.recloser.settings.f51.pickup_a = 190
    ws = Workspace()
    ws.selection.recloser = True
    rep = ReportComposer().compose(p, ws, [StudyKind.RECLOSER])
    keys = [s.key for s in rep.sections]
    assert not [k for k in keys if k.startswith("fd_")]
    assert "prot_recloser" in keys


def test_feeder_dev_titles_validation_and_docx(tmp_path):
    assert is_implemented("feeder_dev")
    assert booklet_title("feeder_dev") == "دفترچه مطالعات توسعه و اصلاح شبکه"
    assert header_title("feeder_dev") == "دفترچه مطالعات توسعه و اصلاح فیدر"
    p = _fd_project()
    items = validate_project(p, AppSettings())
    assert not [i for i in items if i.status == ERROR]
    # خروجی Word: جلد + سربرگ + جدول‌ها
    from app.report import pipeline
    from app.report.template_manager import TemplateManager
    from app.rules.rule_engine import RuleEngine
    ws = Workspace()
    ws.selection.feeder_dev = True
    p._workspace = ws
    rep = pipeline.build_sections(p, AppSettings(), TemplateManager(), RuleEngine(),
                                  tmp_path / "charts")
    assert [s.key for s in rep.sections] == ["fd_intro", "fd_network", "fd_before",
                                             "fd_plan", "fd_after", "fd_conclusion"]
    out = pipeline.generate_docx(p, AppSettings(), rep, tmp_path / "fd.docx")
    from docx import Document
    doc = Document(str(out))
    texts = [par.text for par in doc.paragraphs]
    assert any("دفترچه مطالعات توسعه و اصلاح شبکه" in t for t in texts)
    hdr = doc.sections[-1].header
    htexts = [q.text for q in hdr.paragraphs] + \
        [c.text for t in hdr.tables for row in t.rows for c in row.cells]
    assert any("دفترچه مطالعات توسعه و اصلاح فیدر" in t for t in htexts)
    assert any("FD-15-001" in t for t in htexts)


def test_feeder_dev_page_save_and_load(tmp_path):
    w = _window(tmp_path, recloser=False, feeder_dev=True)
    w.manager.project.report_type = "feeder_dev"
    w.navigate("feeder_dev")
    pg = w.feeder_dev_page
    pg.plan_kind.setCurrentText("اصلاح خط موجود")
    pg.substation.setText("پست خلخال")
    pg.table.insertRow(0)
    pg.table.setItem(0, 0, QTableWidgetItem("پست خلخال — فیدر ۴"))
    pg.table.setItem(0, 1, QTableWidgetItem("20"))
    pg.table.setItem(0, 3, QTableWidgetItem("3.2"))
    pg.min_voltage_kv_before.setText("18.8")
    pg.min_voltage_kv_after.setText("19.6")
    pg.conductor_type.setText("هاینا ۱۲۶")
    w.navigate("project")                        # collect + ترک صفحه
    w.navigate("feeder_dev")                     # بازگشت + refresh
    fd = w.manager.project.feeder_dev
    assert fd.plan_kind == "اصلاح خط موجود"
    assert fd.substation == "پست خلخال"
    assert len(fd.feeders) == 1
    assert fd.feeders[0].voltage_kv == 20 and fd.feeders[0].peak_load_mw == 3.2
    assert fd.min_voltage_kv_before == 18.8 and fd.min_voltage_kv_after == 19.6
    # فرم پس از refresh همان داده‌ها را نشان می‌دهد
    assert pg.table.rowCount() == 1
    assert pg.table.item(0, 0).text() == "پست خلخال — فیدر ۴"


def test_feeder_dev_sections_for_selection_and_tree(tmp_path):
    from app.core.project_tree import build_dashboard, build_project_tree
    from app.core.report_sections import sections_for_selection
    sel = StudySelection(feeder_dev=True)
    keys = [k for k, _ in sections_for_selection(sel)]
    assert keys == ["fd_intro", "fd_network", "fd_before", "fd_plan",
                    "fd_after", "fd_conclusion"]
    # انتخاب حفاظتی بدون توسعه فیدر → هیچ fd_
    keys2 = [k for k, _ in sections_for_selection(StudySelection(recloser=True))]
    assert not [k for k in keys2 if k.startswith("fd_")]
    # درخت و داشبورد
    p = _fd_project()
    ws = Workspace()
    ws.selection.feeder_dev = True
    tree = build_project_tree(p, ws)
    inp = next(n for n in tree.children if n.key == "step.input")
    assert any(n.key == "input.feeder_dev" for n in inp.children)
    dash = build_dashboard(p, ws)
    assert dash["feeder_dev"]["label"] == "توسعه و اصلاح فیدر"
    assert dash["feeder_dev"]["status"] == "complete"
