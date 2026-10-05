# -*- coding: utf-8 -*-
"""v1.1.0 — قابلیت گزارش‌سازی سکشنالایزر (ساختار اولیه؛ جزئیات متعاقباً).

چرخه کامل باید کار کند: پروژه سکشنالایزر → اعتبارسنجی → پیش‌نمایش → تولید Word.
"""
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from docx import Document

from app.core.models import Project, SectionalizerInfo, project_from_dict, to_dict
from app.core.report_types import (REPORT_HEAVY, REPORT_RECLOSER, REPORT_SECTIONALIZER,
                                   booklet_title)
from app.core.settings import AppSettings
from app.core.validation import has_errors, validate_project
from app.report import pipeline
from app.report.template_manager import TemplateManager
from app.rules.rule_engine import RuleEngine

S = AppSettings()


def make_secz_project(**over) -> Project:
    p = Project(name="مطالعه سکشنالایزر فیدر 7 گلخانه",
                report_number="SC-01", date_jalali="1405/07/20",
                report_type=REPORT_SECTIONALIZER,
                sectionalizer=SectionalizerInfo(
                    feeder_name="فیدر 7 گلخانه",
                    location="محل پیشنهادی: تیر شماره 120",
                    purpose="محدودسازی خاموشی هنگام خطای انتهای فیدر"))
    for k, v in over.items():
        setattr(p, k, v)
    return p


def test_validation_passes_without_power_data():
    items = validate_project(make_secz_project(), S)
    assert not has_errors(items)


def test_validation_requires_feeder_name():
    p = make_secz_project()
    p.sectionalizer.feeder_name = ""
    items = validate_project(p, S)
    assert has_errors(items)
    assert any("فیدر هدف" in i.message for i in items)


def test_report_sections_and_titles(tmp_path):
    report, docx, _ = pipeline.generate_report_full(
        make_secz_project(), S, tmp_path / "c", tmp_path,
        TemplateManager(), RuleEngine())
    titles = [s.title for s in report.sections]
    assert titles == ["مقدمه", "مشخصات فیدر و تجهیزات موجود",
                      "محل پیشنهادی نصب سکشنالایزر", "بررسی حفاظتی و هماهنگی حفاظت",
                      "نتیجه‌گیری و پیشنهادات"]
    intro = "\n".join(b.text for b in report.sections[0].paragraphs())
    assert "فیدر 7 گلخانه" in intro and "محدودسازی خاموشی" in intro
    assert "تیر شماره 120" in intro
    # بخش‌های بعدی فعلاً جای‌دار دارند
    assert any("پس از تکمیل اطلاعات" in b.text
               for b in report.sections[1].paragraphs())


def test_docx_generated_with_sectionalizer_booklet(tmp_path):
    report, docx, _ = pipeline.generate_report_full(
        make_secz_project(), S, tmp_path / "c", tmp_path,
        TemplateManager(), RuleEngine())
    doc = Document(str(docx))
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads[0].startswith("1- مقدمه") and len(heads) == 5
    xml = zipfile.ZipFile(docx).read("word/document.xml").decode("utf-8")
    assert "گزارش مطالعات نصب سکشنالایزر" in xml          # جلد + سربرگ
    assert "مطالعات مربوط به مطالعه نصب سکشنالایزر بر روی فیدر 7 گلخانه" in xml


def test_manual_text_edit_works_for_sectionalizer(tmp_path):
    p = make_secz_project()
    p.manual_texts["secz_spec"] = "مشخصات فیدر مطابق اطلاعات امور مربوطه تکمیل شد."
    report, _, _ = pipeline.generate_report_full(p, S, tmp_path / "c", tmp_path,
                                                 TemplateManager(), RuleEngine())
    sec = next(s for s in report.sections if s.key == "secz_spec")
    assert any("تکمیل شد" in b.text for b in sec.paragraphs())
    assert getattr(sec, "manual", False)


def test_serialization_roundtrip():
    p2 = project_from_dict(to_dict(make_secz_project()))
    assert p2.report_type == REPORT_SECTIONALIZER
    assert p2.sectionalizer.feeder_name == "فیدر 7 گلخانه"
    assert p2.title_text() == "مطالعه نصب سکشنالایزر بر روی فیدر 7 گلخانه"


def test_report_type_registry():
    assert booklet_title(REPORT_HEAVY).startswith("دفترچه مطالعات تأمین برق")
    assert "سکشنالایزر" in booklet_title(REPORT_SECTIONALIZER)
    assert "ریکلوزر" in booklet_title(REPORT_RECLOSER)
    # پروژه قدیمی بدون فیلد نوع → متقاضیان سنگین
    p_old = project_from_dict({"name": "قدیمی"})
    assert p_old.is_heavy and p_old.report_type == REPORT_HEAVY
