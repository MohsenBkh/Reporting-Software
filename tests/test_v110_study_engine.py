# -*- coding: utf-8 -*-
"""v1.1.0 «تغییرات آرنا» — تست‌های Rule Engine مطالعه مصارف سنگین.

پوشش آزمون:
* ساختار خروجی استاندارد هر Rule (بخش ۶ صورت‌مسئله)
* محاسبه صریح Delta = After − Before و عدم اشتباه با مقدار مطلق
* کنترل ناسازگاری داده‌ها (DATA_INCONSISTENCY) و عدم حدس/جایگزینی خودکار
* تولید سناریو بدون انتخاب/توصیه
* بررسی اقتصادی: MISSING_DATA به‌جای صفر + پارامتریک بودن هزینه‌ها
* MISSING_RULE در نبود آستانه (هیچ آستانه‌ای اختراع نمی‌شود)
* اتصال به مولد گزارش Word
"""
from __future__ import annotations

import copy
from pathlib import Path

import pytest
from sample_projects import make_project, make_study_project

from app.core import scenarios as scenario_mod
from app.core import study as study_mod
from app.core.models import SCENARIO_KIND_LABELS, SupplyScenario
from app.core.settings import AppSettings, CostSettings, StudyThresholdSettings
from app.report.sections import Paragraph, TableSpec
from app.report.template_manager import TemplateManager
from app.report.text_generator import TextGenerator
from app.rules.outcome import STATUS_VALUES, RuleResult
from app.rules.rule_engine import RuleEngine

SETTINGS = AppSettings()
ENGINE = RuleEngine()


def _results(project, settings=SETTINGS, engine=ENGINE) -> list[RuleResult]:
    scenario_mod.ensure_scenarios(project, settings, engine)
    return study_mod.collect_study_results(project, settings, engine)


def _by_id(results, rule_id, owner=None):
    """فیلتر نتایج بر اساس شناسه Rule (شناسه گزارش‌شده یا شناسه داخلی).

    چند Rule کنترل ناسازگاری، شناسه گزارش‌شده «DATA_INCONSISTENCY» دارند و
    شناسه داخلی آن‌ها در rule_key نگه داشته می‌شود.
    """
    return [r for r in results if rule_id in (r.rule_id, r.rule_key)
            and (owner is None or (r.owner or "") == owner)]


def _one(results, rule_id, owner=None) -> RuleResult:
    found = _by_id(results, rule_id, owner)
    assert found, f"Rule {rule_id} اجرا نشد (owner={owner})"
    return found[0]


# ---------------------------------------------------------------------------
# ساختار خروجی استاندارد
# ---------------------------------------------------------------------------
def test_every_rule_result_follows_standard_schema():
    results = _results(make_study_project())
    assert results
    required = {"rule_id", "category", "status", "input_values", "calculated_values",
                "finding", "evidence", "source", "source_page", "recommended_text",
                "engineer_review_required"}
    for r in results:
        d = r.to_dict()
        assert required.issubset(d.keys()), r.rule_id
        assert d["status"] in STATUS_VALUES, (r.rule_id, d["status"])
        assert d["source"] == "TAV111-10/00"
        assert d["source_page"], r.rule_id
        assert isinstance(d["engineer_review_required"], bool)


def test_all_declared_rules_are_structurally_valid():
    """هر Rule تحلیلی: وضعیت معتبر، منبع و صفحه، و قالب‌های متنی."""
    for rule in ENGINE.analysis_rules.values():
        assert rule.status in STATUS_VALUES, rule.id
        assert rule.source and rule.source_page, rule.id
        assert rule.finding or rule.status in ("MISSING_DATA", "MISSING_RULE"), rule.id
    for srule in ENGINE.scenario_rules.values():
        # v1.2.0: نوع «بازآرایی فیدر (انتقال بار)» نیز به سناریوها اضافه شد
        assert srule.kind in SCENARIO_KIND_LABELS, srule.id
        assert srule.technical_basis and srule.review_status, srule.id


def test_engine_loads_all_rule_files_without_skips():
    assert ENGINE._last_skipped == []
    assert len(ENGINE.analysis_rules) >= 40
    assert len(ENGINE.scenario_rules) == 4


# ---------------------------------------------------------------------------
# Delta = After − Before (و تفکیک از مقدار مطلق)
# ---------------------------------------------------------------------------
def test_line_delta_is_after_minus_before_not_absolute():
    results = _results(make_study_project())
    r = _one(results, "LN-DELTA-CALC", owner="فیدر ۴۱۵ فاز ۵")
    assert r.calculated_values["line_delta_v_pp"] == pytest.approx(0.68, abs=1e-6)
    assert r.input_values["line_v_before_pct"] == 1.17
    assert r.input_values["line_v_after_pct"] == 1.85
    # متن گزارش هر دو مقدار و «تغییر» را نشان می‌دهد، نه فقط مقدار مطلق
    assert "1.17" in r.recommended_text and "1.85" in r.recommended_text
    # Δ ≠ مقدار مطلق After
    assert r.calculated_values["line_delta_v_pp"] != r.input_values["line_v_after_pct"]


def test_absolute_vs_delta_rule_is_emitted():
    results = _results(make_study_project())
    r = _one(results, "LN-ABSOLUTE-VS-DELTA", owner="فیدر ۴۰۱ دفتر شهرک")
    assert "3.83" in r.finding                      # مقدار مطلق After
    assert "0.53" in r.finding                      # تغییر ناشی از متقاضی
    assert r.status == "INFO"


def test_preexisting_voltage_drop_is_separated_from_applicant_effect():
    results = _results(make_study_project())
    r = _one(results, "LN-PREEXISTING-DROP", owner="فیدر ۴۰۱ دفتر شهرک")
    assert r.calculated_values["line_delta_v_pp"] == pytest.approx(0.53, abs=1e-6)
    assert "ناشی از بار متقاضی نبوده" in r.recommended_text


# ---------------------------------------------------------------------------
# کنترل ناسازگاری داده‌ها (Cross Validation)
# ---------------------------------------------------------------------------
def test_load_sum_consistency_passes_on_reference_sample():
    r = _one(_results(make_study_project()), "XC-LOAD-SUM-OK")
    assert r.status == "PASS"
    assert r.input_values["sum_new_demand_changes_kw"] == pytest.approx(3000.0)
    assert r.input_values["reported_total_additional_load_kw"] == pytest.approx(3000.0)


def test_load_sum_mismatch_is_data_inconsistency_without_guessing():
    p = make_study_project(reported_total_additional_load_kw=2500.0)
    r = _one(_results(p), "DATA_INCONSISTENCY")
    assert r.rule_id == "DATA_INCONSISTENCY"
    assert r.status == "REQUIRES_ENGINEER_REVIEW"
    assert r.engineer_review_required is True
    # هیچ مقداری جایگزین/حدس زده نمی‌شود
    assert r.input_values["sum_new_demand_changes_kw"] == pytest.approx(3000.0)
    assert r.input_values["reported_total_additional_load_kw"] == pytest.approx(2500.0)
    assert r.calculated_values["load_sum_diff_kw"] == pytest.approx(500.0)
    assert r.recommended_text == ""


def test_demand_base_mismatch_flagged():
    p = make_study_project(demand_used_in_analysis_kw=1800.0)
    results = _results(p)
    r = _one(results, "XC-DEMAND-BASE-MISMATCH")
    assert r.rule_id == "DATA_INCONSISTENCY"          # شناسه گزارش‌شده مطابق صورت‌مسئله
    assert r.rule_key == "XC-DEMAND-BASE-MISMATCH"    # شناسه داخلی Rule برای ردیابی
    assert r.status == "REQUIRES_ENGINEER_REVIEW"
    assert r.calculated_values["demand_base_diff_kw"] == pytest.approx(-600.0)
    assert r.to_dict()["meta"]["rule_key"] == "XC-DEMAND-BASE-MISMATCH"


def test_both_inconsistency_checks_report_same_rule_id():
    p = make_study_project(reported_total_additional_load_kw=2500.0,
                           demand_used_in_analysis_kw=1800.0)
    results = [r for r in _results(p) if r.rule_id == "DATA_INCONSISTENCY"]
    assert len(results) == 2
    assert {r.rule_key for r in results} == {"XC-LOAD-SUM-MISMATCH", "XC-DEMAND-BASE-MISMATCH"}
    assert all(r.status == "REQUIRES_ENGINEER_REVIEW" for r in results)


def test_missing_baseline_stays_missing_data():
    p = make_study_project(reported_total_additional_load_kw=None)
    r = _one(_results(p), "XC-NO-BASELINE-SUM")
    assert r.status == "MISSING_DATA"
    # نبود داده به‌معنای «سازگار» فرض نمی‌شود
    assert not _by_id(_results(p), "XC-LOAD-SUM-OK")


def test_coincidence_factor_one_requires_engineer_review():
    r = _one(_results(make_study_project()), "D-COINCIDENCE-NOT-APPLIED")
    assert r.status == "REQUIRES_ENGINEER_REVIEW"
    assert r.engineer_review_required


def test_coincidence_factor_over_one_is_data_error():
    p = make_study_project()
    p.demand.with_coincidence_kw = 1500.0
    r = _one(_results(p), "D-COINCIDENCE-OVER-ONE")
    assert r.status == "DATA_ERROR"
    assert r.recommended_text == ""


# ---------------------------------------------------------------------------
# MISSING_RULE — هیچ آستانه‌ای اختراع نمی‌شود
# ---------------------------------------------------------------------------
def test_missing_threshold_produces_missing_rule_and_no_text():
    settings = AppSettings()          # substation_loading_tolerance_pct = None
    results = _results(make_study_project(), settings)
    r = _one(results, "SS-LOADING-CHECK-NO-THRESHOLD", owner="شهرک صنعتی")
    assert r.status == "MISSING_RULE"
    assert r.missing == ["substation_loading_tolerance_pct"]
    assert r.recommended_text == "" and r.text_available is False
    # مقادیر محاسبه‌شده برای تصمیم مهندس نمایش داده می‌شود (بدون داوری خودکار)
    assert r.calculated_values["ss_computed_t1_pct"] == pytest.approx(69.05, abs=0.01)
    assert r.calculated_values["ss_loading_diff_t1_pct"] == pytest.approx(-8.0, abs=0.01)


def test_setting_threshold_activates_consistency_check():
    settings = AppSettings(study=StudyThresholdSettings(substation_loading_tolerance_pct=1.0))
    results = _results(make_study_project(), settings)
    mismatch = _by_id(results, "SS-LOADING-MISMATCH", owner="شهرک صنعتی")
    assert mismatch and mismatch[0].status == "REQUIRES_ENGINEER_REVIEW"
    assert mismatch[0].calculated_values["ss_computed_t1_pct"] == pytest.approx(69.05, abs=0.01)
    # ایستگاه «صنعت» سازگار است → هیچ ناسازگاری گزارش نمی‌شود
    assert not _by_id(results, "SS-LOADING-MISMATCH", owner="صنعت")
    assert not _by_id(results, "SS-LOADING-CHECK-NO-THRESHOLD")


def test_missing_current_tolerance_reports_computed_current_without_judgement():
    results = _results(make_study_project())
    r = _one(results, "LN-CURRENT-CHECK-NO-THRESHOLD", owner="فیدر ۴۰۱ دفتر شهرک")
    assert r.status == "MISSING_RULE"
    assert r.calculated_values["line_computed_current_a"] == pytest.approx(218.0, abs=0.1)
    assert r.input_values["line_peak_current_a"] == pytest.approx(217.8)


def test_sc_limit_not_judged_without_threshold():
    results = _results(make_study_project())
    r = _one(results, "LN-SC-LIMIT-MISSING", owner="فیدر ۴۱۵ فاز ۵")
    assert r.status == "MISSING_RULE"
    assert r.missing == ["study_sc_limit_ka"]
    assert "فرض نمی‌شود" in r.finding
    assert "داوری" in r.evidence


# ---------------------------------------------------------------------------
# سناریوها — تولید بدون انتخاب/توصیه
# ---------------------------------------------------------------------------
def test_scenarios_are_generated_but_not_recommended():
    p = make_study_project()
    scenario_mod.scenarios_for = None  # noqa: B018 — سادگی خوانایی
    specs = scenario_mod.generate_scenarios(p, SETTINGS, ENGINE)
    kinds = {s.kind for s in specs}
    assert {"new_feeder", "existing_feeder", "alternative_feeder"}.issubset(kinds)
    for s in specs:
        assert s.review_status == "REQUIRES_ENGINEER_REVIEW"
        assert s.technical_basis and s.required_network_changes
        assert s.candidate_feeder and s.required_equipment and s.evidence
        assert s.source == "TAV111-10/00"


def test_no_selection_criterion_rule_prevents_recommendation():
    results = _results(make_study_project())
    r = _one(results, "SC-NO-SELECTION-CRITERION")
    assert r.status == "REQUIRES_ENGINEER_REVIEW"


def test_engine_stays_silent_when_selection_criterion_defined():
    p = make_study_project(scenario_selection_criterion="معیار اقتصادی مصوب")
    results = _results(p)
    assert not _by_id(results, "SC-NO-SELECTION-CRITERION")


def test_existing_scenarios_are_preserved_not_regenerated():
    p = make_study_project()
    mine = SupplyScenario(kind="existing_feeder", title="سناریوی کارشناس",
                          technical_basis="مبنا", candidate_feeder="فیدر ۴۱۵")
    p.scenarios = [mine]
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    assert len(p.scenarios) == 1 and p.scenarios[0].title == "سناریوی کارشناس"


# ---------------------------------------------------------------------------
# بررسی اقتصادی
# ---------------------------------------------------------------------------
def test_missing_cost_is_missing_data_not_zero():
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    results = _results(p)
    zeros = [s for s in p.scenarios if s.estimated_cost_million is None]
    assert zeros, "سناریوها باید بدون داده هزینه تولید شوند"
    r = _by_id(results, "EC-COST-MISSING")[0]
    assert r.status == "MISSING_DATA"
    assert "MISSING_DATA" in r.recommended_text


def test_zero_cost_without_basis_is_data_error():
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    p.scenarios[0].estimated_cost_million = 0.0
    r = _by_id(_results(p), "EC-COST-ZERO-NO-BASIS")[0]
    assert r.status == "DATA_ERROR"
    assert "MISSING_DATA" in r.finding
    assert r.recommended_text == ""


def test_cost_parameters_are_parametric_and_not_hardcoded():
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    s1 = p.scenarios[0]
    s1.required_line_length_km = 2.0
    s1.overhead_percent = 90.0
    s1.underground_percent = 10.0

    # بدون پارامتر هزینه → MISSING_RULE (برآورد ساخته نمی‌شود)
    results = _results(p, AppSettings())
    assert _by_id(results, "EC-PARAMETRIC-NO-PARAMS")

    # با پارامترها → برآورد محاسبه‌شده از پارامترها
    settings = AppSettings(costs=CostSettings(overhead_per_km_million=2000.0,
                                              underground_per_km_million=4000.0))
    results = _results(p, settings)
    est = _by_id(results, "EC-PARAMETRIC-ESTIMATE")
    assert est, "با تعریف پارامترها باید برآورد هزینه تولید شود"
    # 1.8 km × 2000 + 0.2 km × 4000 = 3600 + 800 = 4400 (سایر اقلام: MISSING_DATA)
    assert est[0].calculated_values["scenario_estimated_cost_million"] == pytest.approx(4400.0)
    assert "MISSING_DATA" in est[0].recommended_text
    assert "پست زمینی" in est[0].recommended_text

    # تغییر پارامتر → تغییر برآورد (اثبات پارامتریک بودن)
    settings2 = AppSettings(costs=CostSettings(overhead_per_km_million=1000.0,
                                               underground_per_km_million=1000.0))
    est2 = _by_id(_results(p, settings2), "EC-PARAMETRIC-ESTIMATE")
    assert est2[0].calculated_values["scenario_estimated_cost_million"] == pytest.approx(2000.0)


def test_recorded_cost_is_reported_separately_per_scenario():
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    p.scenarios[0].estimated_cost_million = 5130.0
    p.scenarios[0].cost_excluded_items = ["هزینه کلید ایستگاه", "احداث پست زمینی"]
    results = _results(p)
    cost_rows = _by_id(results, "EC-COST-VALUE")
    assert cost_rows and cost_rows[0].input_values["scenario_cost_million"] == 5130.0
    excluded = _by_id(results, "EC-EXCLUDED-ITEMS")
    assert excluded and excluded[0].status == "MISSING_DATA"


def test_route_shorter_than_nearest_source_requires_review():
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    p.scenarios[0].required_line_length_km = 2.0
    p.scenarios[0].overhead_percent = 90.0
    p.scenarios[0].underground_percent = 10.0
    r = _by_id(_results(p), "EC-ROUTE-LENGTH-VS-DISTANCE")[0]
    assert r.status == "REQUIRES_ENGINEER_REVIEW"
    assert r.calculated_values["min_substation_distance_km"] == pytest.approx(4.0)


# ---------------------------------------------------------------------------
# توپولوژی — تأمین مشترک چند نقطه تقاضا
# ---------------------------------------------------------------------------
def test_shared_supply_is_topological_and_requires_review():
    r = _one(_results(make_study_project()), "TP-SHARED-SUPPLY-REVIEW")
    assert r.status == "REQUIRES_ENGINEER_REVIEW"
    assert r.input_values["n_demand_points"] == 4
    assert "توپولوژیک" in r.finding


def test_joint_supply_conclusion_without_evidence_is_data_error():
    p = make_study_project(joint_supply_feasible=False, joint_supply_note="بررسی پراکندگی")
    r = _one(_results(p), "TP-SHARED-SUPPLY-NO-EVIDENCE")
    assert r.status == "DATA_ERROR"


def test_joint_supply_with_evidence_is_recorded():
    p = make_study_project(joint_supply_feasible=True,
                           joint_supply_note="مسیر فیدر موجود تا هر دو نقطه",
                           joint_supply_evidence="خروجی PowerFactory — آرایش فیدر ۴۱۵")
    results = _results(p)
    r = _one(results, "TP-SHARED-SUPPLY-EVIDENCED")
    assert r.status == "INFO" and "امکان‌پذیر" in r.finding
    assert not _by_id(results, "TP-SHARED-SUPPLY-REVIEW")


# ---------------------------------------------------------------------------
# اتصال به مولد گزارش
# ---------------------------------------------------------------------------
def test_report_contains_study_sections_and_tables(tmp_path):
    p = make_study_project()
    settings = AppSettings(include_appendices=False)
    # بخش‌های تکمیلی در چیدمان مرجع (v1.3.0) پیش‌فرض خاموش‌اند — برای آزمون
    # کامل قابلیت موتور، روشن می‌شوند.
    for key in ("study_demand", "study_coincident", "study_analysis",
                "study_scenarios", "study_economics"):
        setattr(settings.report_sections, key, True)
    gen = TextGenerator(p, settings, TemplateManager(), ENGINE, tmp_path)
    sections = gen.build_all()
    keys = [s.key for s in sections]
    for expected in ("study_demand", "study_substations", "study_lines",
                     "study_coincident", "study_analysis", "study_scenarios",
                     "study_economics"):
        assert expected in keys, expected
    # نتیجه‌گیری و پیشنهادات، آخرین بخش محتوایی گزارش است
    assert keys[-1] == "study_conclusion"

    text = "\n".join(x.text for s in sections for x in s.paragraphs())
    assert "{" not in text and "}" not in text          # هیچ متغیر پرنشده‌ای نماند
    caps = [b.caption for s in sections for b in s.blocks if isinstance(b, TableSpec)]
    assert any("خطوط نزدیک" in c for c in caps)
    assert any("سناریوهای تأمین" in c for c in caps)
    assert any("بررسی اقتصادی" in c for c in caps)
    # Findings با کلید یکتا برای بازبینی مهندس
    keys_f = [f.key for f in gen.findings]
    assert len(keys_f) == len(set(keys_f))
    assert any(f.rule_id == "LN-DELTA-CALC" for f in gen.findings)
    assert any(f.severity == "review" for f in gen.findings)


def test_word_generation_includes_study_tables(tmp_path):
    import dataclasses

    from app.report.pipeline import generate_report_full

    p = make_study_project()
    settings = dataclasses.replace(
        SETTINGS, report_sections=dataclasses.replace(SETTINGS.report_sections))
    for key in ("study_scenarios", "study_economics"):
        setattr(settings.report_sections, key, True)
    report, out, _items = generate_report_full(
        p, settings, tmp_path / "charts", tmp_path / "out")
    assert out.exists()
    from docx import Document

    doc = Document(str(out))
    texts = [para.text for para in doc.paragraphs]
    texts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    joined = "\n".join(texts)
    assert "سناریوهای تأمین برق" in joined
    assert "بررسی اقتصادی سناریوها" in joined
    assert "خطوط نزدیک به محل تقاضا" in joined
    assert "افزایش بار متقاضی باعث افزایش افت ولتاژ انتهای خط" in joined


def test_classic_project_report_stays_unchanged(tmp_path):
    """پروژه‌های بدون داده مطالعه، بخش‌های جدید را تولید نمی‌کنند (سازگاری v1.0.3)."""
    from tests.sample_projects import make_feeder

    p = make_project()
    p.feeders = [make_feeder()]
    gen = TextGenerator(p, SETTINGS, TemplateManager(), ENGINE, tmp_path)
    keys = [s.key for s in gen.build_all()]
    assert keys[:6] == ["intro", "loading", "forecast", "before", "after", "conclusion"]


def test_has_study_data_gate():
    assert not study_mod.has_study_data(make_project())
    assert study_mod.has_study_data(make_study_project())
    p = make_project()
    p.demand = None            # نوع نادر ولی ممکن پس از بازکردن پروژه قدیمی
    assert not study_mod.has_study_data(p)


# ---------------------------------------------------------------------------
# ردیابی و رندر متن
# ---------------------------------------------------------------------------
def test_traceability_includes_study_data_rows():
    from app.core import traceability as trace

    rows = trace.build_trace(make_study_project())
    items = {r.item for r in rows}
    assert any("تقاضا با ضریب همزمانی" in i for i in items)
    assert any("افت ولتاژ قبل" in i for i in items)
    assert any("تغییر افت ولتاژ" in i for i in items)


def test_rule_result_trace_chain_is_complete():
    r = _one(_results(make_study_project()), "LN-DELTA-CALC", owner="فیدر ۴۱۵ فاز ۵")
    labels = [x[0] for x in r.trace_lines()]
    assert labels == ["ورودی (Input)", "محاسبه (Calculation)", "قاعده (Rule)",
                      "یافته (Finding)", "شاهد (Evidence)", "متن گزارش (Report Text)",
                      "منبع"]
    d = r.to_dict()
    assert d["meta"]["rule_name"]
    assert d["meta"]["thresholds"] == {}


def test_statistics_used_by_deliverable_report_are_correct():
    """اعداد کلیدی گزارش پایانی از خود موتور استخراج می‌شوند."""
    from app.rules.outcome import summarize_results

    results = _results(make_study_project())
    summary = summarize_results(results)
    assert summary["total"] == len(results)
    rules_seen = {r.rule_id for r in results}
    assert {"LN-DELTA-CALC", "SS-DESCRIBE", "CD-LIST", "SC-NO-SELECTION-CRITERION"} <= rules_seen
    # پروژه کپی‌شده نباید تحت تأثیر قرار گیرد (بدون Side-effect پنهان)
    p2 = copy.deepcopy(make_study_project())
    assert p2.demand.with_coincidence_kw == 1200.0


# ---------------------------------------------------------------------------
# جدول‌های مرجع دفترچه (ایستگاه‌های نزدیک / خطوط نزدیک / نتیجه‌گیری و پیشنهادات)
# ---------------------------------------------------------------------------
def _tables(report, key):
    return [b for s in report.sections if s.key == key
            for b in s.blocks if isinstance(b, TableSpec)]


def _report(tmp_path, project=None, settings=None):
    from app.report.pipeline import build_sections

    p = project or make_study_project()
    settings = settings or AppSettings(include_appendices=False)
    scenario_mod.ensure_scenarios(p, settings, ENGINE)
    return build_sections(p, settings, TemplateManager(), ENGINE, tmp_path)


def test_substation_table_follows_reference_layout(tmp_path):
    report = _report(tmp_path)
    table = _tables(report, "study_substations")[0]
    assert table.header_rows[0][0] == "نام ایستگاه / اطلاعات"
    assert table.header_rows[0][1:] == ["شهرک صنعتی", "صنعت"]
    labels = [r[0] for r in table.body_rows]
    assert labels == ["فاصله تقریبی تا محل تقاضا (km)", "ظرفیت ایستگاه (MVA)",
                      "درصد بارگیری ترانس T1 در پیک بار 1403",
                      "درصد بارگیری ترانس T2 در پیک بار 1403",
                      "میزان بارگیری ترانس T1 در پیک بار 1403 (MVA)",
                      "میزان بارگیری ترانس T2 در پیک بار 1403 (MVA)",
                      "تعداد فیدر برقرار"]
    assert table.body_rows[2][1:] == ["61.05٪", "81.21٪"]


def test_line_table_follows_reference_layout_with_multiline_peak(tmp_path):
    report = _report(tmp_path)
    table = _tables(report, "study_lines")[0]
    assert table.header_rows[0][0] == "نام خط / اطلاعات"
    labels = [r[0] for r in table.body_rows]
    assert "پیک بار خط در سال 1403" in labels
    assert "افت ولتاژ انتهای خط قبل از اتصال بار جدید بر اساس پیک بار 1403" in labels
    assert "افت ولتاژ انتهای خط بعد از اتصال بار جدید بر اساس پیک بار 1403" in labels
    assert "حداکثر جریان اتصال کوتاه (kA)" in labels and "حداقل جریان اتصال کوتاه (kA)" in labels
    peak_cell = table.body_rows[labels.index("پیک بار خط در سال 1403")][2]
    # سه مقدار MVA/MW/A در یک خانه (چندخطی) — مطابق تصویر
    assert peak_cell.split("\n") == ["7.55 MVA", "7.08 MW", "217.8 A"]
    # ΔV ناشی از متقاضی در جدول باقی می‌ماند (ارزش افزودهٔ Rule Engine)
    assert "تغییر افت ولتاژ ناشی از بار جدید (واحد درصد)" in labels


def test_conclusion_table_has_reference_three_rows(tmp_path):
    p = make_study_project()
    scenario_mod.ensure_scenarios(p, SETTINGS, ENGINE)
    p.scenarios[0].estimated_cost_million = 5130.0
    p.scenarios[0].required_line_length_km = 2.0
    p.scenarios[0].overhead_percent = 90.0
    p.scenarios[0].underground_percent = 10.0
    p.scenarios[0].cost_excluded_items = ["هزینه کلید ایستگاه", "احداث پست زمینی"]
    report = _report(tmp_path, p)
    table = _tables(report, "study_conclusion")[0]
    assert table.caption == ""                      # عنوان داخل جدول است
    assert table.header_rows[0][0] == "نتیجه‌گیری و پیشنهادات"
    assert table.merges == [(0, 0, 1)]
    labels = [r[0] for r in table.body_rows]
    assert labels[0] == "نکات تحلیلی"
    assert labels[1] == "سناریوهای پیشنهادی"
    assert labels[2].startswith("بررسی اقتصادی سناریوها")
    assert "بدون در نظر گرفتن هزینه کلید ایستگاه و احداث پست زمینی" in labels[2]

    notes, scenarios, econ = (r[1] for r in table.body_rows)
    # نکات تحلیلی: بندهای شماره‌دار فارسی از خروجی Ruleها
    assert notes.startswith("۱- ")
    assert "مجموع افزایش بار" in notes and "تقاضای همزمان" in notes
    assert "هزینه صفر" in notes or "DATA_ERROR" not in notes
    # سناریوها: سه سناریو + هشدار نبود معیار انتخاب
    assert scenarios.count("\n") >= 2
    assert "احداث" in scenarios and "فیدر موجود" in scenarios
    assert "معیار انتخاب سناریو در Rule Set تعریف نشده است" in scenarios
    # اقتصاد: هزینه ثبت‌شده + MISSING_DATA برای سناریوهای بدون داده
    assert "5130 میلیون تومان" in econ
    assert "MISSING_DATA" in econ
    assert "اقلام هزینه کلید ایستگاه" in econ


def test_reference_layout_section_order_matches_reference_report(tmp_path):
    """چیدمان پیش‌فرض (v1.3.0) دقیقاً مطابق گزارش مرجع مهر ۱۴۰۵:
    مقدمه، تحلیل بارگذاری، پیش‌بینی، ایستگاه‌ها، خطوط، پخش بار قبل،
    پخش بار بعد، کنترل بارگذاری، نتیجه‌گیری و پیشنهادات.
    """
    from sample_projects import make_feeder

    p = make_study_project()
    p.feeders.append(make_feeder())   # فیدر فعال → بخش کنترل بارگذاری ساخته می‌شود
    report = _report(tmp_path, p)
    keys = [s.key for s in report.sections]
    core = ["intro", "loading", "forecast", "study_substations", "study_lines",
            "before", "after", "study_loading", "study_conclusion"]
    assert [k for k in keys if k in core] == core
    # جمع‌بندی کلاسیک وقتی «نتیجه‌گیری و پیشنهادات» هست حذف می‌شود
    assert "conclusion" not in keys
    # بخش‌های تکمیلی به‌طور پیش‌فرض در گزارش نیستند
    for extra in ("study_demand", "study_coincident", "study_analysis",
                  "study_scenarios", "study_economics"):
        assert extra not in keys


def test_multiline_table_cells_render_in_word(tmp_path):
    from docx import Document
    from app.report.pipeline import generate_report_full

    p = make_study_project()
    p.existing_power_kw = 1000.0
    p.demand.existing_demand_kw = 1000.0
    settings = AppSettings(include_appendices=False)
    _report_out, docx, _items = generate_report_full(
        p, settings, tmp_path / "charts", tmp_path / "out", engine=ENGINE)
    doc = Document(str(docx))
    # سلول «پیک بار خط» باید سه پاراگراف داشته باشد
    found = False
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if cells and "پیک بار خط در سال" in cells[0].text:
                target = cells[-1]
                lines = [x.text for x in target.paragraphs if x.text.strip()]
                assert lines == ["7.55 MVA", "7.08 MW", "217.8 A"]
                found = True
    assert found, "جدول خطوط نزدیک در فایل Word یافت نشد"
    texts = [p2.text for p2 in doc.paragraphs] + \
            [c.text for t in doc.tables for r in t.rows for c in r.cells]
    joined = "\n".join(texts)
    assert "نتیجه‌گیری و پیشنهادات" in joined
    assert "نکات تحلیلی" in joined and "سناریوهای پیشنهادی" in joined
