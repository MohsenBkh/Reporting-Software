# -*- coding: utf-8 -*-
"""آزمون‌های v2.5.0 —

۱) کدگذاری مدارک مطابق «کدگذاری مطالعات.xlsx» (DM/PD/FD/AP/CP - کد امور - شماره).
۲) اسکرول صفحه «تصاویر و شکل‌های گزارش» (رفع فشردگی/هم‌پوشانی فرم).
۳) عنوان گزارش از «عنوان پروژه» صفحه اطلاعات پروژه.
۴) نیروگاه‌ها و اتوبوستر در ورودی همه مطالعات.
۵) شینه‌های کاندید در ورودی توسعه و اصلاح فیدر (+ جدول مقایسه در گزارش).
۶) موتور قواعد «توسعه و اصلاح فیدر» بر اساس سه دفترچه نمونه.
۷) واژه‌های انگلیسی متن گزارش با فونت Times New Roman.
"""
from __future__ import annotations

import os
import tempfile

import pytest

os.environ.setdefault("REPORTFORGE_HOME", tempfile.mkdtemp(prefix="rf_home_v250_"))

from app import pyside_compat  # noqa: F401
from PySide6.QtWidgets import QApplication, QScrollArea, QTableWidgetItem, QWidget

QApplication.instance() or QApplication([])

from app.core.models import (AutoBoosterInfo, CandidateBus, FeederDevRow, PowerPlant,
                             Project, project_from_dict, to_dict)
from app.core.report_types import REPORT_TYPE_FEEDER_DEV
from app.core.settings import AppSettings
from app.core.studies import StudyKind
from app.core.workspace import workspace_from_project
from app.report.composer import ReportComposer
from app.report.sections import TableSpec


class _Mgr:
    def __init__(self):
        self.project = Project()
    image_path = staticmethod(lambda i: None)


def _sample_project() -> Project:
    """پروژه معادل دفترچه نمونه «اصلاح خط اوچ‌بلاغ ← کوراماللو» (FD-26-001)."""
    p = Project()
    p.name = "دفترچه مطالعه اصلاح خط مسیر اوچ بلاغ به سمت کوراماللو بر روی فیدر 5 گرمی"
    p.report_number = "FD-26-001"
    p.date_jalali = "1405/08/01"
    p.report_type = REPORT_TYPE_FEEDER_DEV
    fd = p.feeder_dev
    fd.plan_kind = "اصلاح خط موجود"
    fd.substation = "پست فوق توزیع گرمی"
    fd.feeders = [FeederDevRow(name="فیدر 5 گرمی", voltage_kv=20, length_km=395,
                               peak_load_mw=5.5, min_load_mw=2.3,
                               loading_pct=78.5, power_factor=0.92)]
    fd.min_voltage_kv_before = 19.36
    fd.min_voltage_kv_after = 19.36
    fd.losses_kw_before = 1064
    fd.losses_kw_after = 1064
    fd.max_loading_pct_before = 53
    fd.max_loading_pct_after = 53
    fd.candidate_buses = [CandidateBus("روستای کوراماللو", 19.64, 21.14, 23.19, 21.67),
                          CandidateBus("روستای اوچ بولاغ", 19.64, 21.14, 23.37, 21.86)]
    p.has_power_plants = "yes"
    p.power_plants = [PowerPlant("chp", 1.5, None, "پس از انشعاب فرعی اوچ بولاغ"),
                      PowerPlant("solar", 3.0), PowerPlant("solar", 3.0)]
    p.autobooster = AutoBoosterInfo("yes", 1, 200, "ابتدای انشعاب")
    return p


# ---------------------------------------------------------------------------
# ۱) کدگذاری مدارک
# ---------------------------------------------------------------------------
def test_report_coding_structure_matches_xlsx():
    from app.core.report_coding import (OFFICE_CODES, STUDY_CODES, describe_report_number,
                                        is_valid_report_number, make_report_number,
                                        parse_report_number)
    assert STUDY_CODES["feeder_dev"] == "FD"
    assert STUDY_CODES["protective_device"] == "PD"
    assert STUDY_CODES["heavy_applicant"] == "DM"
    assert STUDY_CODES["autobooster"] == "AP"
    assert STUDY_CODES["capacitor"] == "CP"
    # کد امورها مطابق فایل
    assert OFFICE_CODES["برق اردبیل"] == "18" and OFFICE_CODES["مشگین‌شهر"] == "13"
    assert OFFICE_CODES["انگوت"] == "26" and OFFICE_CODES["خلخال"] == "16"
    # مثال خود فایل: PD-13-001
    assert make_report_number("PD", "13", 1) == "PD-13-001"
    assert parse_report_number("PD-13-001") == ("PD", "13", 1)
    assert "مشگین" in describe_report_number("PD-13-001")
    assert is_valid_report_number("FD-26-002")
    assert not is_valid_report_number("FD26001") and not is_valid_report_number("XX-13-001")


def test_report_coding_office_name_tolerance():
    from app.core.report_coding import office_code_for_name, study_code_for_report_type
    assert office_code_for_name("پارس آباد") == "15"
    assert office_code_for_name("امور برق گرمی") == "14"
    assert office_code_for_name("برق اردبیل") == "18"
    assert office_code_for_name("") == ""
    assert study_code_for_report_type("recloser") == "PD"
    assert study_code_for_report_type("feeder_dev") == "FD"
    assert study_code_for_report_type("heavy") == "DM"


def test_project_page_code_builder(monkeypatch):
    from app.ui.project_page import ProjectPage
    pp = ProjectPage(AppSettings())
    pp.load_from(Project())
    i = next(i for i in range(pp.cb_study_code.count()) if pp.cb_study_code.itemData(i) == "FD")
    pp.cb_study_code.setCurrentIndex(i)
    j = next(i for i in range(pp.cb_office_code.count()) if pp.cb_office_code.itemData(i) == "26")
    pp.cb_office_code.setCurrentIndex(j)
    pp.sp_serial.setValue(2)
    pp._insert_code()
    assert pp.ed_number.text() == "FD-26-002"
    assert "انگوت" in pp.lbl_code_hint.text()
    pp.ed_number.setText("هرچی")
    assert "هم‌خوان نیست" in pp.lbl_code_hint.text()


def test_validation_warns_on_nonstandard_report_number():
    from app.core.validation import WARNING, validate_project
    p = _sample_project()
    p.report_number = "گزارش-۱"
    items = validate_project(p, AppSettings())
    assert any(i.status == WARNING and "کدگذاری" in (i.message + i.hint) for i in items)
    p.report_number = "FD-26-001"
    items2 = validate_project(p, AppSettings())
    assert not any("کدگذاری" in (i.message + i.hint) for i in items2 if i.status == WARNING)


# ---------------------------------------------------------------------------
# ۲) اسکرول صفحه تصاویر
# ---------------------------------------------------------------------------
def test_images_page_form_is_scrollable():
    from app.ui.images_page import ImagesPage
    page = ImagesPage(_Mgr())
    scrolls = page.findChildren(QScrollArea)
    assert scrolls, "فرم مشخصات شکل باید داخل QScrollArea باشد"
    assert any(s.widgetResizable() and s.widget() is not None for s in scrolls)


# ---------------------------------------------------------------------------
# ۳) عنوان گزارش از «عنوان پروژه»
# ---------------------------------------------------------------------------
def test_title_text_prefers_project_name():
    p = Project()
    p.report_type = REPORT_TYPE_FEEDER_DEV
    p.name = "دفترچه مطالعه جابه‌جایی خط 20 کیلوولت روستای قره آغاج"
    assert p.title_text() == p.name
    p.name = ""
    assert p.title_text() != ""        # متن خودکار قبلی حفظ می‌شود


def test_word_cover_uses_project_title(tmp_path):
    import docx
    from app.report.word_generator import WordGenerator
    p = _sample_project()
    ws = workspace_from_project(p)
    report = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    out = tmp_path / "fd.docx"
    WordGenerator(p, AppSettings()).generate(report, out)
    text = "\n".join(par.text for par in docx.Document(str(out)).paragraphs)
    assert p.name in text


# ---------------------------------------------------------------------------
# ۴) نیروگاه‌ها و اتوبوستر
# ---------------------------------------------------------------------------
def test_power_plants_and_booster_roundtrip():
    p = _sample_project()
    p2 = project_from_dict(to_dict(p))
    assert p2.has_power_plants == "yes"
    assert [pl.kind for pl in p2.power_plants] == ["chp", "solar", "solar"]
    assert p2.power_plants[0].rated_mw == 1.5
    assert p2.power_plants[0].kind_label == "CHP"
    assert p2.autobooster.present == "yes" and p2.autobooster.current_a == 200


def test_network_assets_widget_save_load():
    from app.ui.shell.network_assets import NetworkAssetsWidget
    mgr = _Mgr()
    w = NetworkAssetsWidget(mgr)
    w.load_from(mgr.project)
    w.cb_has_plants.setCurrentIndex(1)              # وجود دارد
    w.sp_plant_count.setValue(2)
    assert len(w._plant_forms) == 2                 # به تعداد، فرم باز می‌شود
    f0 = w._plant_forms[0]
    k = next(i for i in range(f0.cb_kind.count()) if f0.cb_kind.itemData(i) == "wind")
    f0.cb_kind.setCurrentIndex(k)
    f0.rated_mw.setText("2.5")
    f0.location_note.setText("تپه بادی")
    w.cb_has_booster.setCurrentIndex(1)
    w.booster_current.setText("200")
    w.save_to_project()
    assert mgr.project.has_power_plants == "yes"
    assert mgr.project.power_plants[0].kind == "wind"
    assert mgr.project.power_plants[0].rated_mw == 2.5
    assert mgr.project.autobooster.current_a == 200
    # بارگذاری دوباره در نمونه جدید
    w2 = NetworkAssetsWidget(mgr)
    w2.load_from(mgr.project)
    assert len(w2._plant_forms) == 2
    assert w2._plant_forms[0].cb_kind.currentData() == "wind"
    # «وجود ندارد» فهرست را پاک می‌کند
    w2.cb_has_plants.setCurrentIndex(2)
    w2.save_to_project()
    assert mgr.project.power_plants == []


def test_assets_available_in_all_studies():
    """ویجت نیروگاه/اتوبوستر در هر سه مسیر ورود داده حاضر است."""
    from app.ui.shell.feeder_dev_page import FeederDevPage
    from app.ui.shell.network_assets import NetworkAssetsWidget
    from app.ui.shell.network_editor import NetworkEditor
    from app.ui.shell.study_pages import HeavyStudyPage
    mgr = _Mgr()
    fd = FeederDevPage(mgr)
    assert isinstance(fd.assets, NetworkAssetsWidget)
    ne = NetworkEditor(mgr, lambda: None)
    assert isinstance(ne.assets, NetworkAssetsWidget)
    ne._set_form_enabled(False)                      # بدون فیدر انتخابی
    assert ne.assets.isEnabled()                     # کارت مشترک فعال می‌ماند
    h = HeavyStudyPage(QWidget(), QWidget(), QWidget(), QWidget())
    assert h.count() == 4 and h.tabText(3) == "نیروگاه‌ها و اتوبوستر"


# ---------------------------------------------------------------------------
# ۵) شینه‌های کاندید
# ---------------------------------------------------------------------------
def test_feeder_dev_page_candidate_buses():
    from app.ui.shell.feeder_dev_page import FeederDevPage
    mgr = _Mgr()
    page = FeederDevPage(mgr)
    page.refresh()
    page._add_bus_row()                              # دکمه «+ شینه کاندید»
    page.bus_table.setItem(0, 0, QTableWidgetItem("روستای کوراماللو"))
    page.bus_table.setItem(0, 1, QTableWidgetItem("19.64"))
    page.bus_table.setItem(0, 2, QTableWidgetItem("21.14"))
    page.bus_table.setItem(0, 3, QTableWidgetItem("23.19"))
    page.bus_table.setItem(0, 4, QTableWidgetItem("21.67"))
    page.save_to_project()
    buses = mgr.project.feeder_dev.candidate_buses
    assert len(buses) == 1 and buses[0].name == "روستای کوراماللو"
    assert buses[0].voltage_kv_after == 21.14
    assert buses[0].trafo_loading_pct_before == 23.19
    page2 = FeederDevPage(mgr)
    page2.refresh()
    assert page2.bus_table.rowCount() == 1


def test_composer_bus_comparison_table():
    p = _sample_project()
    ws = workspace_from_project(p)
    report = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    fda = next(s for s in report.sections if s.key == "fd_after")
    tabs = [b for b in fda.blocks if isinstance(b, TableSpec)]
    bus_tab = next(t for t in tabs if "شینه‌های کاندید" in t.caption)
    assert bus_tab.body_rows[0][0] == "روستای کوراماللو"
    assert bus_tab.body_rows[0][2] == "21.14"


# ---------------------------------------------------------------------------
# ۶) موتور قواعد توسعه و اصلاح فیدر
# ---------------------------------------------------------------------------
def test_rules_high_loading_and_seasonal_swing():
    from app.core.feeder_dev_rules import evaluate
    p = _sample_project()
    keys = {f.key.split(".")[1] for f in evaluate(p)}
    assert "loading" in keys          # بارگیری 78.5٪ → نسبتاً بالا (نمونه ۱)
    assert "swing" in keys            # 5.5/2.3 → نوسان فصلی (نمونه ۱/۲)
    assert "pf" in keys               # 0.92 → مناسب (نمونه ۲)


def test_rules_critical_loading_and_low_pf():
    from app.core.feeder_dev_rules import WARNING, evaluate
    p = Project()
    p.feeder_dev.feeders = [FeederDevRow(name="فیدر 1 اجیرلو", loading_pct=95,
                                         power_factor=0.8)]
    fs = evaluate(p)
    assert any(f.severity == WARNING and "ظرفیت کامل" in f.text for f in fs)
    assert any(f.severity == WARNING and "ضریب قدرت" in f.title for f in fs)


def test_rules_plants_and_booster():
    from app.core.feeder_dev_rules import evaluate
    p = _sample_project()
    texts = [f.text for f in evaluate(p)]
    assert any("جریان معکوس" in t and "6" in t for t in texts)   # 3+3 مگاوات خورشیدی
    assert any("CHP" in t for t in texts)
    assert any("اتوبوستر" in t and "200" in t for t in texts)


def test_rules_conclusion_unchanged_matches_sample2():
    from app.core.feeder_dev_rules import conclusion_paragraph
    p = _sample_project()
    text = conclusion_paragraph(p)
    assert "ثابت مانده" in text and "اثر منفی" in text          # ادبیات نمونه ۲
    assert "19.36" in text and "1064" in text and "53" in text
    assert "کوراماللو" in text and "بهبود" in text              # بهبود ولتاژ شینه (نمونه ۱)


def test_rules_losses_reduction_matches_sample3():
    from app.core.feeder_dev_rules import conclusion_paragraph, evaluate
    p = Project()
    fd = p.feeder_dev
    fd.losses_kw_before = 1614
    fd.losses_kw_after = 675
    fs = evaluate(p)
    assert any("939" in f.text and "کاهش" in f.text for f in fs)  # نمونه ۳
    assert "939" in conclusion_paragraph(p)


def test_rules_degradation_is_warning():
    from app.core.feeder_dev_rules import WARNING, evaluate
    p = Project()
    p.feeder_dev.min_voltage_kv_before = 19.5
    p.feeder_dev.min_voltage_kv_after = 18.9
    fs = evaluate(p)
    assert any(f.severity == WARNING and "بدتر شده" in f.text for f in fs)


def test_rules_bus_voltage_limits():
    from app.core.feeder_dev_rules import WARNING, evaluate
    p = Project()
    p.feeder_dev.feeders = [FeederDevRow(name="ف", voltage_kv=20)]
    p.feeder_dev.candidate_buses = [CandidateBus("انتهای فیدر", 16.56, 18.92)]  # 0.828 pu
    fs = evaluate(p)
    assert any(f.severity == WARNING and "پایین‌تر" in f.text for f in fs)     # نمونه ۳


def test_rules_incomplete_data_warning():
    from app.core.feeder_dev_rules import WARNING, evaluate
    p = Project()
    p.feeder_dev.losses_kw_before = 1000                      # «بعد» ثبت نشده
    fs = evaluate(p)
    assert any(f.severity == WARNING and "ناقص" in f.title for f in fs)


def test_composer_network_section_has_rule_paragraphs_and_plant_table():
    p = _sample_project()
    ws = workspace_from_project(p)
    report = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    fdn = next(s for s in report.sections if s.key == "fd_network")
    texts = [getattr(b, "text", "") for b in fdn.blocks]
    assert any("چشمگیر" in t for t in texts)                 # تحلیل بارگیری
    assert any("جریان معکوس" in t for t in texts)
    tabs = [b for b in fdn.blocks if isinstance(b, TableSpec)]
    plant_tab = next(t for t in tabs if "نیروگاه" in t.caption)
    assert plant_tab.body_rows[0][0] == "CHP"


def test_validation_surfaces_rule_warnings():
    from app.core.validation import WARNING, validate_project
    p = _sample_project()
    p.feeder_dev.min_voltage_kv_after = None                  # داده ناقص
    items = validate_project(p, AppSettings())
    assert any(i.status == WARNING and "ناقص" in i.message for i in items)


# ---------------------------------------------------------------------------
# ۷) فونت Times New Roman برای واژه‌های انگلیسی
# ---------------------------------------------------------------------------
def test_latin_segments_regex():
    from app.report.word_generator import _LATIN_SEG
    txt = "مطابق دستورالعمل WI150900 و نرم‌افزار DIgSILENT PowerFactory با هادی ACSR-70 mm2"
    found = [m.group() for m in _LATIN_SEG.finditer(txt)]
    assert "WI150900" in found
    assert "DIgSILENT PowerFactory" in found
    assert any(s.startswith("ACSR-70") for s in found)


def test_add_par_splits_latin_runs_without_rtl():
    import docx
    from docx.oxml.ns import qn
    from app.report.word_generator import add_par
    d = docx.Document()
    par = add_par(d, "خروجی نرم‌افزار DIgSILENT PowerFactory مبنای تحلیل است.",
                  font_fa="B Nazanin", size=12)
    latin = [r for r in par.runs if "DIgSILENT" in (r.text or "")]
    assert latin, "قطعه لاتین باید run جداگانه باشد"
    rPr = latin[0]._r.find(qn("w:rPr"))
    rF = rPr.find(qn("w:rFonts"))
    assert rF.get(qn("w:ascii")) == "Times New Roman"
    assert rF.get(qn("w:hAnsi")) == "Times New Roman"
    assert rPr.find(qn("w:rtl")) is None                     # بدون پرچم rtl
    fa = [r for r in par.runs if "خروجی" in (r.text or "")]
    assert fa and fa[0]._r.find(qn("w:rPr")).find(qn("w:rtl")) is not None


def test_word_output_latin_words_in_times_new_roman(tmp_path):
    import docx
    from docx.oxml.ns import qn
    from app.report.word_generator import WordGenerator
    p = _sample_project()
    p.feeder_dev.before_description = "نتایج با نرم‌افزار DIgSILENT PowerFactory محاسبه شد."
    ws = workspace_from_project(p)
    report = ReportComposer().compose(p, ws, [StudyKind.FEEDER_DEV])
    out = tmp_path / "fd2.docx"
    WordGenerator(p, AppSettings()).generate(report, out)
    found = False
    for par in docx.Document(str(out)).paragraphs:
        for r in par.runs:
            if "DIgSILENT" in (r.text or ""):
                rPr = r._r.find(qn("w:rPr"))
                rF = rPr.find(qn("w:rFonts"))
                assert rF.get(qn("w:ascii")) == "Times New Roman"
                assert rPr.find(qn("w:rtl")) is None
                found = True
    assert found
