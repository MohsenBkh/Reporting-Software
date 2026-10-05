# -*- coding: utf-8 -*-
"""v1.3.0 — چیدمان گزارش مطابق گزارش مرجع مهر ۱۴۰۵ + ثبت «نوع گزارش».

گزارش مرجع: «مطالعه درخواست دیماند 1000kW متقاضی گلخانه آقای مجید دهقانی
دهنوی» — ۹ بخش:
    ۱ مقدمه — ۲ تحلیل بارگذاری — ۳ پیش‌بینی — ۴ ایستگاه‌های نزدیک —
    ۵ خطوط نزدیک — ۶ پخش بار قبل — ۷ پخش بار بعد — ۸ کنترل بارگذاری —
    ۹ نتیجه‌گیری و پیشنهادات
"""
from __future__ import annotations

from sample_projects import make_feeder, make_project, make_study_project

from app.core.report_sections import REPORT_SECTIONS, SECTION_ORDER
from app.core.report_types import (REPORT_TYPE_BOOKLETS, REPORT_TYPE_HEAVY,
                                   REPORT_TYPE_RECLOSER,
                                   REPORT_TYPE_SECTIONALIZER, booklet_title,
                                   is_implemented, normalize,
                                   pending_message, report_type_label)
from app.core.settings import AppSettings
from app.report.pipeline import build_sections, report_to_html
from app.report.template_manager import TemplateManager
from app.rules.rule_engine import RuleEngine

ENGINE = RuleEngine()


def _heavy_report(tmp_path, project=None, settings=None):
    p = project or make_study_project()
    p.feeders.append(make_feeder())
    settings = settings or AppSettings(include_appendices=False)
    return build_sections(p, settings, TemplateManager(), ENGINE, tmp_path), p


# ---------------------------------------------------------------------------
# ثبت نوع گزارش (report_types)
# ---------------------------------------------------------------------------
def test_report_type_registry_labels_and_booklets():
    assert report_type_label(REPORT_TYPE_HEAVY) == "متقاضیان سنگین (یک مگاوات و بالاتر)"
    assert report_type_label(REPORT_TYPE_SECTIONALIZER) == "مطالعه سکشنالایزر"
    assert report_type_label(REPORT_TYPE_RECLOSER) == "مطالعه ریکلوزر"
    assert booklet_title(REPORT_TYPE_HEAVY).startswith("دفترچه مطالعات تأمین برق")
    assert booklet_title(REPORT_TYPE_SECTIONALIZER) == "گزارش مطالعه نصب سکشنالایزر"
    assert booklet_title(REPORT_TYPE_RECLOSER) == "گزارش مطالعه نصب ریکلوزر"
    # کلید خالی/ناشناخته → پیش‌فرض متقاضیان سنگین
    assert normalize("") == REPORT_TYPE_HEAVY
    assert normalize("unknown") == REPORT_TYPE_HEAVY
    assert booklet_title("") == REPORT_TYPE_BOOKLETS[REPORT_TYPE_HEAVY]


def test_recloser_is_planned_but_not_implemented():
    assert is_implemented(REPORT_TYPE_HEAVY)
    assert is_implemented(REPORT_TYPE_SECTIONALIZER)
    assert not is_implemented(REPORT_TYPE_RECLOSER)
    msg = pending_message(REPORT_TYPE_RECLOSER)
    assert "ریکلوزر" in msg and "نسخه‌های بعدی" in msg


# ---------------------------------------------------------------------------
# چیدمان بخش‌ها مطابق گزارش مرجع
# ---------------------------------------------------------------------------
def test_section_order_registry_matches_reference_report():
    core = [k for k, _label in REPORT_SECTIONS if k in (
        "intro", "loading", "forecast", "study_substations", "study_lines",
        "before", "after", "study_loading", "study_conclusion")]
    assert core == ["intro", "loading", "forecast", "study_substations",
                    "study_lines", "before", "after", "study_loading",
                    "study_conclusion"]
    assert SECTION_ORDER.index("study_substations") < SECTION_ORDER.index("before")
    assert SECTION_ORDER.index("study_lines") < SECTION_ORDER.index("before")
    assert SECTION_ORDER.index("after") < SECTION_ORDER.index("study_loading")


def test_reference_titles_match_reference_report():
    labels = dict(REPORT_SECTIONS)
    assert labels["study_substations"] == "ایستگاه‌های نزدیک به محل تقاضا"
    assert labels["study_lines"] == "خطوط نزدیک به محل تقاضا"
    assert labels["study_loading"] == "کنترل بارگذاری فیدر و راهکارهای تعدیل بار"
    assert labels["study_conclusion"] == "نتیجه‌گیری و پیشنهادات"


def test_heavy_report_default_layout_is_reference_nine_sections(tmp_path):
    report, _p = _heavy_report(tmp_path)
    keys = [s.key for s in report.sections]
    core = ["intro", "loading", "forecast", "study_substations", "study_lines",
            "before", "after", "study_loading", "study_conclusion"]
    assert [k for k in keys if k in core] == core
    # بخش‌های تکمیلی به‌طور پیش‌فرض خاموش‌اند و جمع‌بندی حذف شده است
    for extra in ("study_demand", "study_coincident", "study_analysis",
                  "study_scenarios", "study_economics", "conclusion"):
        assert extra not in keys


def test_intro_contains_demand_clause_like_reference(tmp_path):
    """بند الف مقدمه (تقاضای درخواستی/ضریب همزمانی) مطابق گزارش مرجع."""
    report, p = _heavy_report(tmp_path)
    intro = next(s for s in report.sections if s.key == "intro")
    text = "\n".join(b.text for b in intro.paragraphs())
    assert "تقاضای درخواستی" in text
    assert "ضریب همزمانی" in text
    assert "نوع متقاضی: افزایش قدرت" in text
    assert "{" not in text and "}" not in text


def test_intro_demand_clause_absent_without_study_data(tmp_path):
    p = make_project()
    p.feeders.append(make_feeder())
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), ENGINE, tmp_path)
    intro = next(s for s in report.sections if s.key == "intro")
    text = "\n".join(b.text for b in intro.paragraphs())
    assert "ضریب همزمانی" not in text


def test_classic_project_keeps_generic_conclusion(tmp_path):
    """پروژه بدون دادهٔ مطالعه: جمع‌بندی کلاسیک همچنان بخش پایانی است."""
    p = make_project()
    p.feeders.append(make_feeder())
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), ENGINE, tmp_path)
    keys = [s.key for s in report.sections]
    assert keys == ["intro", "loading", "forecast", "before", "after",
                    "conclusion"]


def test_extras_can_be_enabled_and_stay_before_conclusion(tmp_path):
    p = make_study_project()
    p.feeders.append(make_feeder())
    settings = AppSettings(include_appendices=False)
    for key in ("study_demand", "study_coincident", "study_analysis",
                "study_scenarios", "study_economics"):
        setattr(settings.report_sections, key, True)
    report = build_sections(p, settings, TemplateManager(), ENGINE, tmp_path)
    keys = [s.key for s in report.sections]
    assert keys[-1] == "study_conclusion"
    for extra in ("study_demand", "study_coincident", "study_analysis",
                  "study_scenarios", "study_economics"):
        assert extra in keys
        assert keys.index(extra) < keys.index("study_conclusion")
        assert keys.index("study_loading") < keys.index(extra)


def test_preview_html_uses_booklet_title(tmp_path):
    report, p = _heavy_report(tmp_path)
    html = report_to_html(report, p, AppSettings())
    assert "دفترچه مطالعات تأمین برق به متقاضیان یک مگاوات و بالاتر" in html
    assert "9- نتیجه‌گیری و پیشنهادات" in html


def test_word_document_uses_booklet_and_reference_sections(tmp_path):
    from docx import Document

    from app.report.pipeline import generate_report_full

    p = make_study_project()
    p.feeders.append(make_feeder())
    settings = AppSettings(include_appendices=False)
    _report, out, _items = generate_report_full(
        p, settings, tmp_path / "charts", tmp_path / "out", engine=ENGINE)
    doc = Document(str(out))
    assert doc.core_properties.subject.startswith("دفترچه مطالعات تأمین برق")
    headings = [para.text for para in doc.paragraphs if para.style.name.startswith("Heading")]
    joined = "\n".join(headings)
    assert "ایستگاه‌های نزدیک به محل تقاضا" in joined
    assert "خطوط نزدیک به محل تقاضا" in joined
    assert "کنترل بارگذاری فیدر و راهکارهای تعدیل بار" in joined
    assert "نتیجه‌گیری و پیشنهادات" in joined
    # ترتیب سرفصل‌ها مطابق گزارش مرجع
    pos = [joined.find(t) for t in ("ایستگاه‌های نزدیک به محل تقاضا",
                                    "خطوط نزدیک به محل تقاضا",
                                    "کنترل بارگذاری فیدر و راهکارهای تعدیل بار",
                                    "نتیجه‌گیری و پیشنهادات")]
    assert all(x >= 0 for x in pos) and pos == sorted(pos)
