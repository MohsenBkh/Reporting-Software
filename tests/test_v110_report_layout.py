# -*- coding: utf-8 -*-
"""v1.1.0 — چیدمان گزارش مطابق گزارش مرجع (مهر ۱۴۰۵، متقاضی مجید دهقانی دهنوی).

ترتیب بخش‌ها:
    1 مقدمه / 2 تحلیل بارگذاری / 3 پیش‌بینی / 4 ایستگاه‌های نزدیک / 5 خطوط نزدیک
    / 6 پخش بار قبل / 7 پخش بار بعد / 8 کنترل بارگذاری / 9 نتیجه‌گیری و پیشنهادات
"""
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from docx import Document

from app.core.models import (AdjacentFeeder, Feeder, NearbyLine, NearbyStation,
                             PowerFlowResult, Project, REQUEST_NEW, project_from_dict,
                             to_dict)
from app.core.settings import AppSettings
from app.report import pipeline
from app.report.sections import TableSpec
from app.report.template_manager import TemplateManager
from app.rules.rule_engine import RuleEngine

S = AppSettings()
# گزارش مرجع بدون پیوست است (تنظیم پیش‌فرض کاربر نیز همین است)
S.include_appendices = False


def make_reference_project() -> Project:
    """بازسازی داده‌های گزارش مرجع: دیماند 1000kW گلخانه آقای دهقانی دهنوی."""
    p = Project(
        name="مطالعه درخواست دیماند 1000kW متقاضی گلخانه آقای مجید دهقانی دهنوی",
        report_number="DM-15-001", date_jalali="1405/07/10",
        applicant_name="آقای مجید دهقانی دهنوی",
        request_type=REQUEST_NEW, requested_power_kw=1000,
        substation="پارس آباد", office="پارس آباد", expert_name="محسن بخشی",
        location_note="38: X: 732021, Y: 4377580",
        location_distance_m=5, conductor_type="AL-126 (ACSR-Hyena)",
    )
    f = Feeder(name="فیدر 7 گلخانه", substation="پارس آباد",
               peak_load_mw=7.12, peak_year=1405, power_factor=0.92,
               peak_current_a=222, capacity_mw=7,
               before=PowerFlowResult(222, 414, 0.904, 0.98),
               after=PowerFlowResult(250, 448, 0.904, 0.972))
    f.adjacent_feeders = [AdjacentFeeder("فیدر 4 اصلاندوز", 3),
                          AdjacentFeeder("فیدر 4 تهیه", 6),
                          AdjacentFeeder("فیدر 2 شهرک", 9.2)]
    p.feeders = [f]
    p.nearby_stations = [NearbyStation(
        name="پارس آباد", distance_km=15, capacity_mva=45,
        t1_loading_pct=70, t2_loading_pct=70,
        t1_loading_mva=11, t2_loading_mva=22, feeder_count=12)]
    p.nearby_lines = [NearbyLine(
        name="7 گلخانه", distance_m=10, peak_mva=7.69, peak_mw=7.12, peak_a=222,
        vdrop_before_pct=0.9, vdrop_after_pct=0.9)]
    p.conclusion_scenarios = ("در این خصوص سناریو بازآرایی فیدر 7 گلخانه با فیدر 4 اصلاندوز "
                              "به عنوان راهکار پیشنهادی مطرح می‌گردد.")
    return p


EXPECTED_SECTIONS = [
    "مقدمه",
    "تحلیل وضعیت بارگذاری فیدر 7 گلخانه",
    "پیش‌بینی پیک بار در افق پنج‌ساله",
    "ایستگاه‌های نزدیک به محل تقاضا",
    "خطوط نزدیک به محل تقاضا",
    "نتایج پخش بار قبل از اتصال بار جدید",
    "نتایج پخش بار پس از اتصال بار جدید",
    "کنترل بارگذاری فیدر و راهکارهای تعدیل بار",
    "نتیجه‌گیری و پیشنهادات",
]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("v110")
    report, docx, items = pipeline.generate_report_full(
        make_reference_project(), S, out / "charts", out,
        TemplateManager(), RuleEngine())
    return report, Path(docx)


# ---------------------------------------------------------------------------
def test_section_order_matches_reference(built):
    report, _ = built
    titles = [s.title for s in report.sections]
    assert titles == EXPECTED_SECTIONS


def test_table_numbering_matches_reference(built):
    report, _ = built
    caps = [b.caption for s in report.sections for b in s.blocks
            if isinstance(b, TableSpec) and b.caption]
    assert caps[0].startswith("جدول 1: وضعیت بارگذاری")
    assert caps[1] == "جدول 2: ایستگاه‌های نزدیک به محل تقاضا"
    assert caps[2] == "جدول 3: خطوط نزدیک به محل تقاضا"
    assert caps[3].startswith("جدول 4: نتایج")
    assert caps[4] == "جدول 5: کنترل بارگذاری کل فیدر و راهکار پیشنهادی"


def test_stations_table_content(built):
    report, _ = built
    sec = next(s for s in report.sections if s.key == "stations")
    tbl = next(b for b in sec.blocks if isinstance(b, TableSpec))
    assert tbl.header_rows[0] == ["نام ایستگاه / اطلاعات", "پارس آباد"]
    flat = [c for row in tbl.body_rows for c in row]
    assert "فاصله تقریبی تا محل تقاضا (km)" in flat
    assert "15" in flat and "45" in flat and "70٪" in flat and "12" in flat
    assert any("در پیک بار 1405" in c for c in flat)
    # پاراگراف مقدمه با خلاصه ایستگاه
    intro = sec.paragraphs()[0].text
    assert "ایستگاه «پارس آباد»" in intro and "15" in intro


def test_lines_table_and_note(built):
    report, _ = built
    sec = next(s for s in report.sections if s.key == "lines")
    tbl = next(b for b in sec.blocks if isinstance(b, TableSpec))
    flat = [c for row in tbl.body_rows for c in row]
    # پیک بار خط: سه مقدار در یک سلول چندخطی
    peak = next(c for c in flat if "MVA" in c)
    assert "7.69 MVA" in peak and "7.12 MW" in peak and "222 A" in peak
    # تغییر افت ولتاژ با علامت — مطابق گزارش مرجع «+0.00»
    assert "+0.00" in flat
    # یادداشت «افت ولتاژ موجود ناشی از متقاضی نیست»
    notes = [b.text for b in sec.paragraphs() if b.style == "note"]
    assert any("ناشی از بار متقاضی نبوده" in t for t in notes)


def test_after_section_has_bullets_and_result_table(built):
    report, _ = built
    sec = next(s for s in report.sections if s.key == "after")
    bullets = [b.text for b in sec.paragraphs() if b.style == "bullet"]
    assert any("از 222 آمپر به 250 آمپر" in t for t in bullets)
    assert any("از 414 کیلووات به 448 کیلووات" in t for t in bullets)
    assert any("از 0.98 به 0.972 پریونیت" in t for t in bullets)
    tbl = next(b for b in sec.blocks if isinstance(b, TableSpec))
    flat = [c for row in tbl.body_rows for c in row]
    assert "250" in flat and "448" in flat and "0.904" in flat


def test_control_section_critical_with_adjacent_solution(built):
    report, _ = built
    sec = next(s for s in report.sections if s.key == "control")
    text = "\n".join(b.text for b in sec.paragraphs())
    # 7.12 + 1 = 8.12 مگاوات → 116٪ معیار → شدیداً بحرانی
    assert "شدیداً بحرانی" in text and "8.12" in text
    assert "فیدر همجوار" in text and "فیدر 4 اصلاندوز" in text
    tbl = next(b for b in sec.blocks if isinstance(b, TableSpec))
    row = tbl.body_rows[0]
    assert row[0] == "فیدر 7 گلخانه" and row[1] == "7.12" and row[2] == "1"
    assert row[3] == "8.12" and row[4] == "شدیداً بحرانی"
    assert "فیدر 4 اصلاندوز با بار 3 مگاوات" in row[5]
    assert "فیدر 2 شهرک با بار 9.2 مگاوات" in row[5]


def test_control_section_omitted_when_no_need(tmp_path):
    """پروژه سالم (بار کل در محدوده معیار) → بخش کنترل حذف می‌شود."""
    p = Project(name="ت", report_number="DM-1", date_jalali="1405/01/01",
                applicant_name="م", request_type=REQUEST_NEW, requested_power_kw=100)
    p.feeders = [Feeder(name="فیدر 1", peak_load_mw=2.0, peak_year=1405, capacity_mw=7,
                        before=PowerFlowResult(60, 50, 0.97, 0.99),
                        after=PowerFlowResult(65, 55, 0.969, 0.989))]
    report, _, _ = pipeline.generate_report_full(p, S, tmp_path / "c", tmp_path)
    keys = [s.key for s in report.sections]
    assert "control" not in keys and "stations" not in keys and "lines" not in keys


def test_conclusion_is_table_with_analysis_and_scenarios(built):
    report, _ = built
    sec = next(s for s in report.sections if s.key == "conclusion")
    assert sec.title == "نتیجه‌گیری و پیشنهادات"
    tbl = next(b for b in sec.blocks if isinstance(b, TableSpec))
    assert tbl.caption == "" and tbl.merges == [(0, 0, 1)]
    assert tbl.header_rows[0][0] == "نتیجه‌گیری و پیشنهادات"
    assert tbl.body_rows[0][0] == "نکات تحلیلی"
    assert tbl.body_rows[1][0] == "سناریوهای پیشنهادی"
    # متن قاعده‌محور (افت ولتاژ موجود) در سلول نکات تحلیلی
    assert "ناشی از بار متقاضی نیست" in tbl.body_rows[0][1]
    assert "بازآرایی فیدر 7 گلخانه" in tbl.body_rows[1][1]


def test_intro_location_mentions_conductor(built):
    report, _ = built
    intro = report.sections[0]
    text = "\n".join(b.text for b in intro.paragraphs())
    assert "هادی شبکه از نوع AL-126 (ACSR-Hyena)" in text
    assert "تأمین برق به آقای مجید دهقانی دهنوی به ظرفیت 1000 کیلووات" in text


def test_docx_headings_and_cover(built):
    report, path = built
    doc = Document(str(path))
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads == [f"{i + 1}- {t}" for i, t in enumerate(EXPECTED_SECTIONS)]
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    assert "مطالعات مربوط به تأمین برق متقاضی آقای مجید دهقانی دهنوی" in xml
    assert "دفترچه مطالعات تأمین برق به متقاضیان یک مگاوات و بالاتر" in xml
    # آیتم‌های نشانه‌دار بخش 7 در سند درج شده‌اند
    assert "•" in xml or "●" in xml


def test_serialization_roundtrip_new_fields():
    p = make_reference_project()
    p2 = project_from_dict(to_dict(p))
    assert p2.nearby_stations[0].t2_loading_mva == 22
    assert p2.nearby_lines[0].vdrop_delta_pct() == pytest.approx(0.0)
    assert p2.feeders[0].adjacent_feeders[2].load_mw == 9.2
    assert p2.conductor_type == "AL-126 (ACSR-Hyena)"
    assert p2.conclusion_scenarios and p2.report_type == "heavy"
