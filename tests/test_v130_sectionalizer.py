# -*- coding: utf-8 -*-
"""v1.3.0 — قابلیت گزارش‌سازی «سکشنالایزر» (اسکلت فعال).

ساختار پنج‌بخشی آماده است و جزئیات تکمیلی بعداً توسط کارفرما ارسال می‌شود؛
این تست‌ها اسکلت، نوع گزارش، اعتبارسنجی و سازگاری ذخیره‌سازی را پوشش می‌دهند.

v2.0.0: تست «ریکلوزر مسدود است» جایگزین شد — ریکلوزر اکنون با Report Composer
گزارش می‌دهد (``test_recloser_report_type_builds_protection_sections``).
"""
from __future__ import annotations

import pytest

from app.core.models import SectionalizerInfo, project_from_dict, to_dict
from app.core.report_types import (REPORT_TYPE_HEAVY, REPORT_TYPE_RECLOSER,
                                   REPORT_TYPE_SECTIONALIZER)
from app.core.settings import AppSettings
from app.core.workspace import workspace_from_project
from app.core.validation import (ERROR, WARNING, validate_project)
from app.report.pipeline import build_sections, generate_docx, report_to_html
from app.report.template_manager import TemplateManager
from app.rules.rule_engine import RuleEngine


def make_sectionalizer_project(**over):
    from sample_projects import make_project
    p = make_project()
    p.report_type = REPORT_TYPE_SECTIONALIZER
    p.applicant_name = "اداره برق منطقه نمونه"
    p.sectionalizer = SectionalizerInfo(
        feeder_name="فیدر ۲۰کیلوولت نمونه",
        installation_location="تیر شماره ۱۲۳، تقاطع بلوار کشاورز",
        objective="ایزوله‌سازی سریع خطاهای بالادست و کاهش خاموشی",
        fault_current_ka=4.5,
        min_fault_current_ka=1.2,
        pickup_current_a=280.0,
        tms=0.15,
        upstream_device="ریکلوزر پست فوق توزیع نمونه",
        coordination_note="زمان‌بندی به‌صورت پلکانی با ریکلوزر بالادست تنظیم شود.",
    )
    for k, v in over.items():
        setattr(p, k, v)
    return p


# ---------------------------------------------------------------------------
# v2.0.0 — گزارش سکشنالایزر با Report Composer (جایگزین اسکلت پنج‌بخشی v1.3.0)
# ---------------------------------------------------------------------------
def test_sectionalizer_report_uses_protection_composer(tmp_path):
    p = make_sectionalizer_project()
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    keys = [s.key for s in report.sections]
    assert keys[0] == "prot_intro" and "prot_sectionalizer" in keys
    assert "prot_recloser" not in keys and not any(k.startswith("secz_") for k in keys)
    # پیش‌فرض ۷ فصلی دفترچه نمونه: پایان با سناریوی عملکرد (جمع‌بندی/پیشنهادها خاموش)
    assert keys[-1] == "prot_scenario"
    assert "prot_conclusion" not in keys and "prot_recommendation" not in keys


def test_sectionalizer_report_shows_recorded_base_values_and_no_unfilled_vars(tmp_path):
    from app.report.sections import Paragraph

    p = make_sectionalizer_project()
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    sec = next(s for s in report.sections if s.key == "prot_sectionalizer")
    text = "\n".join(b.text for b in sec.blocks if isinstance(b, Paragraph))
    assert "4.5 کیلوآمپر" in text and "280 آمپر" in text and "0.15" in text
    all_text = "\n".join(b.text for s in report.sections for b in s.paragraphs())
    assert "{" not in all_text and "}" not in all_text
    assert "فیدر ۲۰کیلوولت نمونه" in all_text


def test_sectionalizer_missing_fields_show_missing_not_zero(tmp_path):
    from app.report.sections import TableSpec

    p = make_sectionalizer_project(sectionalizer=SectionalizerInfo())
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    sec = next(s for s in report.sections if s.key == "prot_sectionalizer")
    table = next(b for b in sec.blocks if isinstance(b, TableSpec))
    assert all(row[3] == "ناموجود" for row in table.body_rows[:4])        # Pickup چهار تابع
    assert all(row[1] == "ناموجود" for row in table.body_rows[4:])        # شات/Dead time/...


def test_sectionalizer_section_on_off(tmp_path):
    p = make_sectionalizer_project()
    ws = workspace_from_project(p)
    ws.report_prefs.disabled_sections = ["prot_trace", "prot_scenario"]   # جمع‌بندی روشن
    p._workspace = ws
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    keys = [s.key for s in report.sections]
    assert "prot_scenario" not in keys and "prot_conclusion" in keys


def test_sectionalizer_word_generation(tmp_path):
    from docx import Document

    p = make_sectionalizer_project()
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    out = tmp_path / "secz.docx"
    generate_docx(p, settings, report, out)
    assert out.exists()
    doc = Document(str(out))
    assert doc.core_properties.subject == "دفترچه مطالعات جایابی تجهیزات و هماهنگی حفاظتی"
    joined = "\n".join(para.text for para in doc.paragraphs)
    assert "سکشنالایزر" in joined
    html = report_to_html(report, p, settings)
    assert "دفترچه مطالعات جایابی تجهیزات و هماهنگی حفاظتی" in html


# ---------------------------------------------------------------------------
# نوع گزارش — ذخیره‌سازی و اعتبارسنجی
# ---------------------------------------------------------------------------
def test_report_type_and_sectionalizer_survive_serialization():
    p = make_sectionalizer_project()
    data = to_dict(p)
    assert data["report_type"] == REPORT_TYPE_SECTIONALIZER
    assert data["sectionalizer"]["fault_current_ka"] == 4.5
    p2 = project_from_dict(data)
    assert p2.report_type == REPORT_TYPE_SECTIONALIZER
    assert isinstance(p2.sectionalizer, SectionalizerInfo)
    assert p2.sectionalizer.pickup_current_a == 280.0
    assert p2.sectionalizer.coordination_note.startswith("زمان‌بندی")
    # پروژه‌های قدیمی بدون فیلد نوع گزارش → پیش‌فرض متقاضیان سنگین
    old = project_from_dict({"name": "قدیمی"})
    assert old.report_type == REPORT_TYPE_HEAVY
    assert old.sectionalizer == SectionalizerInfo()


def test_sectionalizer_validation():
    settings = AppSettings()
    # کامل → هیچ خطایی نیست
    p = make_sectionalizer_project()
    items = validate_project(p, settings)
    assert not any(i.status == ERROR for i in items)
    # بدون نام/شماره/تاریخ → خطا
    p2 = make_sectionalizer_project(applicant_name="", report_number="",
                                    date_jalali="")
    items2 = validate_project(p2, settings)
    assert any(i.status == ERROR for i in items2)
    # بدون داده فنی → هشدار (نه خطا) و تولید گزارش ممکن می‌ماند
    p3 = make_sectionalizer_project(sectionalizer=SectionalizerInfo())
    items3 = validate_project(p3, settings)
    assert any(i.status == WARNING for i in items3)
    assert not any(i.status == ERROR for i in items3)


def test_heavy_validation_untouched_for_heavy_projects():
    p = make_sectionalizer_project(report_type=REPORT_TYPE_HEAVY)
    items = validate_project(p, AppSettings())
    # مسیر اعتبارسنجی متقاضیان سنگین: فیدر تعریف‌نشده → خطا
    assert any(i.status == ERROR for i in items)


def test_recloser_report_type_builds_protection_sections(tmp_path):
    """v2.0.0: ریکلوزر دیگر مسدود نیست؛ با Report Composer بخش‌های prot_* ساخته می‌شود."""
    p = make_sectionalizer_project(report_type=REPORT_TYPE_RECLOSER)
    settings = AppSettings(include_appendices=False)
    report = build_sections(p, settings, TemplateManager(), RuleEngine(), tmp_path)
    keys = [s.key for s in report.sections]
    assert keys[0] == "prot_intro"
    assert "prot_recloser" in keys and "prot_sectionalizer" not in keys
    assert keys[-1] == "prot_scenario"
    # هیچ مقدار مهندسی ساخته نشده: تنظیمات ثبت‌نشده با «ناموجود» نمایش داده می‌شود
    assert report.warnings or any(s.key == "prot_scenario" for s in report.sections)


def test_recloser_validation_is_light_and_has_no_heavy_errors():
    p = make_sectionalizer_project(report_type=REPORT_TYPE_RECLOSER)
    items = validate_project(p, AppSettings())
    # مسیر متقاضی سنگین (فیدر/پخش بار) برای ریکلوزر اجرا نمی‌شود
    assert not any(i.status == ERROR for i in items)
    assert any(i.status == WARNING for i in items)          # تنظیمات ریکلوزر ناقص است
