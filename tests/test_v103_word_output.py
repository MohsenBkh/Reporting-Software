# -*- coding: utf-8 -*-
"""v1.0.3 — خروجی Word: ساختار، RTL، Caption، TOC، فیلدها، Golden-style Regression."""
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from sample_projects import make_feeder, make_project

from app.core.settings import AppSettings
from app.report import pipeline

S = AppSettings()


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("w")
    p = make_project()
    p.feeders = [make_feeder(years=[1401, 1402, 1403, 1404, 1405], values=[4.0, 4.3, 4.5, 4.9, 5.1])]
    report, docx, items = pipeline.generate_report_full(p, S, out / "charts", out)
    return report, Path(docx)


def _xml(path):
    return zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")


def test_docx_opens_and_has_all_sections(built):
    report, path = built
    doc = Document(str(path))
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(heads) == len(report.sections) >= 6
    assert heads[0].startswith("1-")
    assert any(h.startswith("پیوست") and "-" in h and not h[0].isdigit() for h in heads)


def test_rtl_and_cs_fonts_everywhere(built):
    xml = _xml(built[1])
    assert xml.count("<w:bidi") >= 20 and "w:bidiVisual" in xml
    assert 'w:cs="' in xml


def test_toc_prefilled_and_update_fields_on_open(built):
    z = zipfile.ZipFile(built[1])
    assert "w:updateFields" in z.read("word/settings.xml").decode("utf-8")
    xml = z.read("word/document.xml").decode("utf-8")
    toc = xml[xml.index("TOC \\o"):]
    assert "مقدمه" in toc[:4000]


def test_footer_has_page_x_of_y_and_cover_is_unnumbered(built):
    z = zipfile.ZipFile(built[1])
    footers = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/footer"))
    assert "PAGE" in footers and "NUMPAGES" in footers
    assert "w:titlePg" in z.read("word/document.xml").decode("utf-8")


def test_captions_use_caption_style_and_sequential_numbering(built):
    doc = Document(str(built[1]))
    caps = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    tbl = [int(m.group(1)) for c in caps if (m := re.match(r"جدول (\d+):", c))]
    fig = [int(m.group(1)) for c in caps if (m := re.match(r"شکل (\d+):", c))]
    assert tbl == list(range(1, len(tbl) + 1)) and fig == list(range(1, len(fig) + 1))
    assert tbl and fig


def test_table_header_rows_repeat_and_do_not_split(built):
    xml = _xml(built[1])
    assert "w:tblHeader" in xml and "w:cantSplit" in xml


def test_core_properties(built):
    cp = Document(str(built[1])).core_properties
    assert cp.language == "fa-IR" and "ReportForge" in cp.keywords and cp.title


def test_no_leftover_placeholders_or_none_in_document(built):
    text = "\n".join(p.text for p in Document(str(built[1])).paragraphs)
    assert "{" not in text and "None" not in text and "nan" not in text.lower().split()


def test_golden_structure_regression(built):
    """ساختار طلایی گزارش (ترتیب بخش‌ها و کلیدها) نباید بی‌خبر تغییر کند."""
    keys = [s.key for s in built[0].sections]
    assert keys[:6] == ["intro", "loading", "forecast", "before", "after", "conclusion"]
    assert keys[6:] == ["appendix_trace"]


def test_appendices_can_be_disabled(tmp_path):
    s = AppSettings(); s.include_appendices = False
    p = make_project(); p.feeders = [make_feeder()]
    rep, _d, _i = pipeline.generate_report_full(p, s, tmp_path / "c", tmp_path)
    assert not any(x.key.startswith("appendix") for x in rep.sections)


def test_regeneration_is_deterministic(tmp_path):
    p = make_project(); p.feeders = [make_feeder()]
    t = lambda: "\n".join(x.text for s in pipeline.build_sections(
        p, S, __import__("app.report.template_manager", fromlist=["x"]).TemplateManager(),
        __import__("app.rules.rule_engine", fromlist=["x"]).RuleEngine(), tmp_path / "c").sections
        for x in s.paragraphs())
    assert t() == t()
