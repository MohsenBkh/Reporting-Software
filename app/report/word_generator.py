# -*- coding: utf-8 -*-
"""مولد فایل Word (DOCX) — فارسی، RTL، سربرگ/پاورقی، فهرست مطالب، جداول و شکل‌ها.

از python-docx + تنظیمات مستقیم OOXML برای راست‌به‌چپ استفاده می‌کند.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.core.models import Project, REQUEST_INCREASE
from app import APP_VERSION
from app.core.settings import AppSettings
from app.report.sections import FigureBlock, GeneratedReport, Paragraph, ReportSection, TableSpec
from app.utils.jalali import jalali_month_year
from app.utils.resources import resource_path

HEADER_GRAY = "D9E2F3"
BORDER_COLOR = "7F7F7F"


def default_logo_path() -> Path | None:
    """لوگوی پیش‌فرض همراه برنامه (resource/logo.png) در صورت وجود."""
    p = resource_path("logo.png")
    return p if p.exists() else None


# ---------------------------------------------------------------------------
# ابزارهای OOXML
# ---------------------------------------------------------------------------
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def set_rtl(paragraph, rtl: bool = True) -> None:
    """bidi را در جایگاه درست سند OOXML (قبل از jc) درج می‌کند — idempotent."""
    pPr = paragraph._p.get_or_add_pPr()
    if rtl:
        if pPr.find(qn("w:bidi")) is not None:
            return
        bidi = _el("w:bidi")
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            jc.addprevious(bidi)
        else:
            pPr.append(bidi)


def style_run(run, font_fa: str, size: int, bold: bool = False,
              color: RGBColor | None = None, font_latin: str = "Times New Roman") -> None:
    run.font.name = font_latin
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _el("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:cs"), font_fa)
    # ترتیب فرزندان rPr مطابق اسکیمای OOXML: ... b, bCs, ... sz, szCs, ... rtl
    if rPr.find(qn("w:szCs")) is None:
        szCs = _el("w:szCs", val=str(int(size * 2)))
        sz = rPr.find(qn("w:sz"))
        if sz is not None:
            sz.addnext(szCs)
        else:
            rPr.append(szCs)
    if bold and rPr.find(qn("w:bCs")) is None:
        bCs = _el("w:bCs")
        b = rPr.find(qn("w:b"))
        if b is not None:
            b.addnext(bCs)
        else:
            szCs = rPr.find(qn("w:szCs"))
            if szCs is not None:
                szCs.addprevious(bCs)
            else:
                rPr.append(bCs)
    if rPr.find(qn("w:rtl")) is None:
        rPr.append(_el("w:rtl"))


def add_par(doc_or_cell, text: str, *, font_fa: str, size: int,
            bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            color: RGBColor | None = None, rtl: bool = True,
            space_after: int = 6, font_latin: str = "Times New Roman",
            line_spacing: float = 1.15):
    """افزودن پاراگراف فارسی RTL با استایل کامل."""
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    set_rtl(p, rtl)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    run = p.add_run(text)
    style_run(run, font_fa, size, bold=bold, color=color, font_latin=font_latin)
    return p


def add_table_bidi(table) -> None:
    """جدول راست‌به‌چپ + خطوط کامل — با رعایت ترتیب اسکیمای tblPr."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr.find(qn("w:bidiVisual")) is None:
        bidi = _el("w:bidiVisual")
        # bidiVisual باید قبل از tblW/tblBorders بیاید — بعد از tblStyle
        tbl_style = tblPr.find(qn("w:tblStyle"))
        if tbl_style is not None:
            tbl_style.addnext(bidi)
        else:
            tblPr.insert(0, bidi)
    if tblPr.find(qn("w:tblBorders")) is None:
        borders = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_el(f"w:{side}", val="single", sz="4",
                               space="0", color=BORDER_COLOR))
        tblPr.append(borders)


def shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el("w:shd", val="clear", color="auto", fill=hex_color))


def cell_text(cell, text: str, *, font_fa: str, size: int, bold=False,
              align=WD_ALIGN_PARAGRAPH.CENTER, color: RGBColor | None = None,
              font_latin: str = "Times New Roman") -> None:
    """متن سلول — از «\\n» برای چند خط داخل سلول پشتیبانی می‌کند (مثل گزارش مرجع)."""
    lines = str(text).split("\n")
    cell.text = ""
    _cell_par(cell.paragraphs[0], lines[0], font_fa=font_fa, size=size, bold=bold,
              align=align, color=color, font_latin=font_latin)
    for line in lines[1:]:
        _cell_par(cell.add_paragraph(), line, font_fa=font_fa, size=size, bold=bold,
                  align=align, color=color, font_latin=font_latin)


def _cell_par(p, text: str, *, font_fa: str, size: int, bold=False,
              align=WD_ALIGN_PARAGRAPH.CENTER, color: RGBColor | None = None,
              font_latin: str = "Times New Roman") -> None:
    p.alignment = align
    set_rtl(p)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    style_run(run, font_fa, size, bold=bold, color=color, font_latin=font_latin)


def _field(paragraph, instr: str) -> None:
    """درج فیلد Word (مثل TOC یا PAGE)."""
    r1 = paragraph.add_run()
    fld = _el("w:fldChar", fldCharType="begin")
    r1._r.append(fld)
    r2 = paragraph.add_run()
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    r2._r.append(instrText)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", fldCharType="separate"))
    r4 = paragraph.add_run("…")
    style_run(r4, "B Nazanin", 12)
    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", fldCharType="end"))


# ---------------------------------------------------------------------------
# مولد اصلی
# ---------------------------------------------------------------------------
class WordGenerator:
    def __init__(self, project: Project, settings: AppSettings) -> None:
        self.project = project
        self.settings = settings
        self.f = settings.fonts

    # ------------------------------------------------------------------
    def generate(self, report: GeneratedReport, out_path: str | Path) -> Path:
        doc = Document()
        self._setup_styles(doc)
        self._setup_page(doc)
        self._cover(doc)
        self._toc(doc, [self._heading_text(sec, i + 1) for i, sec in enumerate(report.sections)])
        for i, sec in enumerate(report.sections):
            self._section(doc, sec, number=i + 1)
        self._set_properties(doc)
        self._update_fields_on_open(doc)
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    @staticmethod
    def _is_appendix(sec: ReportSection) -> bool:
        return sec.key.startswith("appendix")

    def _heading_text(self, sec: ReportSection, number: int) -> str:
        return sec.title if self._is_appendix(sec) else f"{number}- {sec.title}"

    def _setup_styles(self, doc: Document) -> None:
        """فونت فارسی و زبان bidi در استایل‌های پایه (Normal / Heading / Caption)."""
        f = self.f
        for name in ("Normal", "Heading 1", "Heading 2", "Caption"):
            try:
                st = doc.styles[name]
            except KeyError:
                continue
            rPr = st.element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = _el("w:rFonts")
                rPr.insert(0, rFonts)
            fa = f.heading_font if name.startswith("Heading") else f.body_font
            for attr in ("ascii", "hAnsi", "eastAsia"):
                rFonts.set(qn(f"w:{attr}"), f.latin_font)
            rFonts.set(qn("w:cs"), fa)
            if rPr.find(qn("w:lang")) is None:
                rPr.append(_el("w:lang", val="en-US", bidi="fa-IR"))
        h1 = doc.styles["Heading 1"]
        h1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        h1.paragraph_format.keep_with_next = True
        h1.paragraph_format.space_before = Pt(12)

    def _set_properties(self, doc: Document) -> None:
        cp = doc.core_properties
        cp.title = self.project.title_text()
        cp.subject = self.project.booklet_title()
        cp.author = self.project.expert_name or self.settings.company_name
        cp.language = "fa-IR"
        cp.keywords = f"ReportForge {APP_VERSION}"

    @staticmethod
    def _update_fields_on_open(doc: Document) -> None:
        """Word هنگام بازشدن فهرست مطالب و شماره صفحات را به‌روز می‌کند."""
        settings = doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            settings.append(_el("w:updateFields", val="true"))

    # ------------------------------------------------------------------
    def _setup_page(self, doc: Document) -> None:
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        self._header_footer(sec)

    def _header_footer(self, sec) -> None:
        f = self.f
        p = self.project
        # --- سربرگ صفحات ۲ به بعد: «دفترچه…» راست | «شماره گزارش / تاریخ» چپ + خط ممتد ---
        header = sec.header
        hp0 = header.paragraphs[0]
        # جدول نامرئی دوستونه (RTL): ستون اول = راست (عنوان دفترچه)، ستون دوم = چپ (شماره/تاریخ)
        ht = header.add_table(rows=1, cols=2, width=Cm(17.0))
        ht.alignment = WD_TABLE_ALIGNMENT.CENTER
        ht.autofit = False
        tblPr = ht._tbl.tblPr
        tblPr.append(_el("w:bidiVisual"))
        # چیدمان ثابت + عرض کل، تا Word عرض ستون‌ها را بازتنظیم نکند
        tblPr.append(_el("w:tblLayout", type="fixed"))
        tblPr.append(_el("w:tblW", w="9639", type="dxa"))
        # بدون حاشیه (جدول نامرئی)
        borders = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_el(f"w:{side}", val="none", sz="0", space="0", color="auto"))
        tblPr.append(borders)
        right_cell, left_cell = ht.rows[0].cells  # در bidiVisual: اولین سلول سمت راست نمایش داده می‌شود
        right_cell.width = Cm(9.5)
        left_cell.width = Cm(7.5)
        # حذف پاراگراف خالی پیش‌فرض بالای جدول (ارتفاع اضافی سربرگ نمی‌گیرد)
        hp0._p.getparent().remove(hp0._p)
        cp = right_cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(cp)
        r = cp.add_run(p.booklet_title())
        style_run(r, f.body_font, 10, bold=True)
        cp2 = left_cell.paragraphs[0]
        cp2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_rtl(cp2)
        r = cp2.add_run(f"شماره گزارش: {p.report_number}        تاریخ: {jalali_month_year(p.date_jalali)}")
        style_run(r, f.body_font, 10, bold=True)
        # خط ممتد زیر سربرگ (پاراگراف جدا با حاشیه پایین)
        hline = header.add_paragraph()
        hline.paragraph_format.space_before = Pt(0)
        hline.paragraph_format.space_after = Pt(0)
        pPr = hline._p.get_or_add_pPr()
        pBdr = _el("w:pBdr")
        pBdr.append(_el("w:bottom", val="single", sz="8", space="1",
                        color=BORDER_COLOR))
        pPr.append(pBdr)
        # --- پاورقی: «صفحه X از Y» ---
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_rtl(fp)
        r = fp.add_run("صفحه ")
        style_run(r, f.body_font, 10)
        self._simple_field(fp, " PAGE ", f.body_font)
        r = fp.add_run(" از ")
        style_run(r, f.body_font, 10)
        self._simple_field(fp, " NUMPAGES ", f.body_font)
        # صفحه جلد بدون سربرگ/پاورقی
        sec.different_first_page_header_footer = True

    @staticmethod
    def _simple_field(paragraph, instr: str, font: str) -> None:
        r1 = paragraph.add_run()
        style_run(r1, font, 10)
        r1._r.append(_el("w:fldChar", fldCharType="begin"))
        r2 = paragraph.add_run()
        style_run(r2, font, 10)
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2._r.append(it)
        r3 = paragraph.add_run()
        style_run(r3, font, 10)
        r3._r.append(_el("w:fldChar", fldCharType="separate"))
        r4 = paragraph.add_run("1")
        style_run(r4, font, 10)
        r5 = paragraph.add_run()
        style_run(r5, font, 10)
        r5._r.append(_el("w:fldChar", fldCharType="end"))

    def _logo_path(self) -> Path | None:
        """لوگوی جلد: اولویت با تنظیم کاربر، سپس لوگوی همراه برنامه."""
        custom = self.settings.company_logo
        if custom:
            p = Path(custom)
            if p.exists():
                return p
        return default_logo_path()

    # ------------------------------------------------------------------
    def _cover(self, doc: Document) -> None:
        f = self.f
        p = self.project
        # لوگوی شرکت بالای نام شرکت
        logo = self._logo_path()
        if logo is not None:
            try:
                lp = doc.add_paragraph()
                lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                lp.paragraph_format.space_after = Pt(8)
                lp.add_run().add_picture(str(logo), height=Cm(3.2))
            except Exception:  # noqa: BLE001 — لوگوی خراب نباید گزارش را متوقف کند
                pass
        add_par(doc, self.settings.company_name, font_fa=f.heading_font,
                size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        add_par(doc, self.settings.unit_name, font_fa=f.body_font,
                size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        add_par(doc, self.settings.office_name, font_fa=f.body_font,
                size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

        # چیدمان جلد مطابق گزارش مرجع: عنوان دفترچه ← موضوع مطالعه ← ماه/سال
        # ← شماره گزارش ← کارشناس مطالعات
        add_par(doc, p.booklet_title(),
                font_fa=f.heading_font, size=22, bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30,
                color=RGBColor(0x1F, 0x4E, 0x79))

        add_par(doc, f"مطالعات مربوط به {p.title_text()}", font_fa=f.heading_font,
                size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)

        add_par(doc, jalali_month_year(p.date_jalali), font_fa=f.heading_font,
                size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=60)

        add_par(doc, f"شماره گزارش: {p.report_number}", font_fa=f.body_font,
                size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        if p.expert_name:
            add_par(doc, f"کارشناس مطالعات: {p.expert_name}", font_fa=f.body_font,
                    size=12, align=WD_ALIGN_PARAGRAPH.CENTER)

        pb = doc.add_paragraph()
        pb.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------
    def _toc(self, doc: Document, titles: list[str]) -> None:
        f = self.f
        add_par(doc, "فهرست مطالب", font_fa=f.heading_font, size=f.heading1_size,
                bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
        p = doc.add_paragraph()
        set_rtl(p)
        # فیلد TOC با نتیجه اولیه (عناوین) — Word هنگام بازشدن شماره صفحات را کامل می‌کند
        p.add_run()._r.append(_el("w:fldChar", fldCharType="begin"))
        r2 = p.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = ' TOC \\o "1-2" \\h \\z \\u '
        r2._r.append(it)
        p.add_run()._r.append(_el("w:fldChar", fldCharType="separate"))
        for i, t in enumerate(titles):
            r = p.add_run(t)
            style_run(r, f.body_font, f.body_size)
            if i < len(titles) - 1:
                r.add_break()
        p.add_run()._r.append(_el("w:fldChar", fldCharType="end"))
        pb = doc.add_paragraph()
        pb.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------
    def _section(self, doc: Document, sec: ReportSection, number: int) -> None:
        f = self.f
        # عنوان با استایل Heading 1 — تا فیلد TOC پس از F9 در Word پر شود
        hp = doc.add_paragraph(style="Heading 1")
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(hp)
        hp.paragraph_format.space_after = Pt(10)
        run = hp.add_run(self._heading_text(sec, number))
        style_run(run, f.heading_font, f.heading1_size, bold=True,
                  color=RGBColor(0x1F, 0x4E, 0x79), font_latin=f.latin_font)
        for block in sec.blocks:
            if isinstance(block, Paragraph):
                if block.style == "bullet":
                    self._bullet(doc, block.text)
                else:
                    style = ("item" if block.style == "item" else "body")
                    add_par(doc, block.text, font_fa=f.body_font, size=f.body_size,
                            bold=(style == "item"),
                            align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif isinstance(block, TableSpec):
                self._table(doc, block)
            elif isinstance(block, FigureBlock):
                self._figure(doc, block)

    def _bullet(self, doc: Document, text: str) -> None:
        """آیتهای نشانه‌دار (مانند «جریان ابتدای فیدر ...») — مطابق استایل لیست گزارش مرجع."""
        f = self.f
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_rtl(p)
        pf = p.paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = 1.15
        pf.right_indent = Cm(0.9)
        run = p.add_run("•  ")
        style_run(run, f.body_font, f.body_size, bold=True, font_latin=f.latin_font)
        run = p.add_run(text)
        style_run(run, f.body_font, f.body_size, font_latin=f.latin_font)

    # ------------------------------------------------------------------
    def _table(self, doc: Document, spec: TableSpec) -> None:
        f = self.f
        # عنوان جدول — بالای جدول (مطابق نمونه‌ها). عنوان خالی = بدون کپشن
        # (مثل جدول نتیجه‌گیری و پیشنهادات که سربرگ ادغام‌شده خودش عنوان است).
        if spec.caption:
            cap = add_par(doc, spec.caption, font_fa=f.body_font, size=f.body_size - 1,
                          bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4,
                          color=RGBColor(0, 0, 0))
            cap.style = doc.styles["Caption"]
            cap.paragraph_format.keep_with_next = True

        n_cols = max((len(r) for r in spec.header_rows + spec.body_rows), default=1)
        n_head = len(spec.header_rows)
        table = doc.add_table(rows=n_head + len(spec.body_rows), cols=n_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_table_bidi(table)

        for ri, row in enumerate(spec.header_rows):
            for ci in range(n_cols):
                text = row[ci] if ci < len(row) else ""
                cell_text(table.cell(ri, ci), text, font_fa=f.body_font,
                          size=f.body_size - 1, bold=True)
                shade_cell(table.cell(ri, ci), HEADER_GRAY)

        for ri, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(_el("w:cantSplit"))
            if ri < n_head:
                trPr.append(_el("w:tblHeader"))
        for (ri, c0, c1) in spec.merges:
            if c1 > c0 and c1 < n_cols:
                # سلول‌های ادغام‌شونده (جز اولی) خالی شوند تا متن سربرگ تکرار نشود
                for ci in range(c0 + 1, c1 + 1):
                    for par in table.cell(ri, ci).paragraphs:
                        for run in list(par.runs):
                            run._r.getparent().remove(run._r)
                merged = table.cell(ri, c0).merge(table.cell(ri, c1))
                # حذف پاراگراف‌های خالی باقی‌مانده از ادغام (سلول سربرگ یک‌خطی بماند)
                for extra in merged.paragraphs[1:]:
                    if not extra.text.strip():
                        extra._p.getparent().remove(extra._p)

        for ri, row in enumerate(spec.body_rows):
            for ci in range(n_cols):
                text = row[ci] if ci < len(row) else ""
                cell = table.cell(n_head + ri, ci)
                cell_text(cell, text, font_fa=f.body_font, size=f.body_size - 1,
                          bold=(ci == 0))
                if ci == 0:
                    shade_cell(cell, HEADER_GRAY)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------
    def _figure(self, doc: Document, fig: FigureBlock) -> None:
        f = self.f
        path = Path(fig.path)
        if not path.exists():
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            run = p.add_run()
            run.add_picture(str(path), width=Cm(15.0))
        except Exception:  # noqa: BLE001 — تصویر خراب نباید کل گزارش را متوقف کند
            return
        p.paragraph_format.keep_with_next = True
        cap = add_par(doc, fig.caption, font_fa=f.body_font, size=f.body_size - 1,
                      bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10,
                      color=RGBColor(0, 0, 0))
        cap.style = doc.styles["Caption"]
